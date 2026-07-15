# ============================================================================
# DISABLED: PILOT-native PPT Copilot implementation (pre-GD-template).
# Superseded by the upstream GD-template implementation from
# syugesh/pilot-voice-agent-backend (feature/ppt-copilot @ 869dd08d9), pasted
# in active below. Kept here commented out for reference only — do not import.
# ============================================================================
# # import asyncio
# # import logging
# 
# # from backend.core.slide_store import _current_slide, _slide_store, resolve_sid
# # from backend.queues.bus import bus
# 
# # """PPT tools — navigate + jump to slide by number or title + summarize."""
# 
# # logger = logging.getLogger("pilot.tools.ppt")
# 
# 
# 
# # async def ppt_navigate(args: dict, session_id: str) -> dict:
# #     direction = args.get("direction","next")
# #     effective_sid = resolve_sid(session_id)
# #     slides = _slide_store.get(effective_sid, [])
# #     total = len(slides)
# #     current = _current_slide.get(effective_sid, 0)
# 
# #     if direction == "next":
# #         if total > 0 and current >= total - 1:
# #             return {
# #                 "spoken_reply": f"You've reached the last slide — slide {total} of {total}. That's the end of the presentation."
# #             }
# #         _current_slide[effective_sid] = min(current + 1, max(total - 1, 0))
# #     elif direction == "prev":
# #         if current <= 0:
# #             return {"spoken_reply": "You're already on the first slide."}
# #         _current_slide[effective_sid] = max(current - 1, 0)
# #     elif direction == "first":
# #         _current_slide[effective_sid] = 0
# #     elif direction == "last":
# #         _current_slide[effective_sid] = max(total - 1, 0)
# 
# #     await bus.emit_event("ppt_command", {"action": direction}, session_id)
# #     return {"status": "ok", "direction": direction, "index": _current_slide.get(effective_sid, 0)}
# 
# 
# # async def ppt_jump_to_title(args: dict, session_id: str) -> dict:
# #     query = args.get("query", "")
# #     slide_number = args.get("slide_number")  # already 0-indexed if from keyword fallback
# #     import re
# 
# #     from backend.queues.bus import bus
# 
# #     effective_sid = resolve_sid(session_id)
# #     slides = _slide_store.get(effective_sid, [])
# 
# #     # If explicit slide number given, use directly
# #     if slide_number is not None:
# #         idx = int(slide_number)
# #         _current_slide[effective_sid] = idx
# #         await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
# #         return {"status": "ok", "index": idx, "title": f"Slide {idx + 1}"}
# 
# #     q = query.lower()
# 
# #     # Numeric match in query — only allow within actual deck bounds
# #     m = re.search(r"\b(\d+)\b", q)
# #     if m:
# #         idx = int(m.group(1)) - 1
# #         if 0 <= idx < len(slides):
# #             _current_slide[effective_sid] = idx
# #             await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
# #             return {"status": "ok", "index": idx}
# 
# #     # Title fuzzy match
# #     best_idx, best_score = 0, 0
# #     for s in slides:
# #         score = sum(1 for w in q.split() if w in s.get("title", "").lower())
# #         if score > best_score:
# #             best_score = score
# #             best_idx = s["index"]
# 
# #     _current_slide[effective_sid] = best_idx
# #     await bus.emit_event("ppt_command", {"action": "goto", "index": best_idx}, session_id)
# #     return {"status": "ok", "index": best_idx}
# 
# 
# # async def ppt_delete_slide(args: dict, session_id: str) -> dict:
# #     from backend.api.ppt import _current_slide, _slide_store, resolve_sid, save_and_render_pptx
# #     from backend.queues.bus import bus
# 
# #     effective_sid = resolve_sid(session_id)
# #     slides = list(_slide_store.get(effective_sid, []))
# #     if not slides:
# #         return {"status": "error", "spoken_reply": "There are no slides in the presentation to delete."}
# 
# #     current_idx = _current_slide.get(effective_sid, 0)
# #     if current_idx < 0 or current_idx >= len(slides):
# #         current_idx = 0
# 
# #     slides.pop(current_idx)
# 
# #     # Re-build slide list preserving images
# #     slides_to_save = []
# #     for s in slides:
# #         slides_to_save.append(
# #             {
# #                 "title": s.get("title", ""),
# #                 "bullets": s.get("bullets", []),
# #                 "notes": s.get("notes", ""),
# #                 "images": s.get("images", []),
# #             }
# #         )
# 
# #     try:
# #         rendered = save_and_render_pptx(slides_to_save, effective_sid, topic="Presentation")
# #         new_idx = min(current_idx, max(len(rendered) - 1, 0))
# #         _current_slide[effective_sid] = new_idx
# 
# #         # Emit reload event
# #         await bus.emit_event(
# #             "ppt_command",
# #             {
# #                 "action": "reload",
# #                 "slides": rendered,
# #                 "filename": f"Presentation ({len(rendered)} slides)",
# #                 "index": new_idx,
# #                 "preserveCurrent": False,
# #             },
# #             session_id,
# #         )
# 
# #         return {"status": "ok", "spoken_reply": "I have deleted the current slide."}
# #     except Exception as e:
# #         return {"status": "error", "spoken_reply": f"Failed to delete slide: {e}"}
# 
# 
# # async def ppt_create_slides(args: dict, session_id: str) -> dict:
# #     import re
# 
# #     from backend.api.ppt import create_presentation_from_prompt
# #     from backend.queues.bus import bus
# 
# #     query = args.get("query", "presentation")
# 
# #     # ── DYNAMIC SLIDE COUNT PARSING ──
# #     # Default is 5 slides
# #     slide_count = 5
# 
# #     word_to_num = {
# #         "one": 1,
# #         "two": 2,
# #         "three": 3,
# #         "four": 4,
# #         "five": 5,
# #         "six": 6,
# #         "seven": 7,
# #         "eight": 8,
# #         "nine": 9,
# #         "ten": 10,
# #         "eleven": 11,
# #         "twelve": 12,
# #         "thirteen": 13,
# #         "fourteen": 14,
# #         "fifteen": 15,
# #         "sixteen": 16,
# #         "seventeen": 17,
# #         "eighteen": 18,
# #         "nineteen": 19,
# #         "twenty": 20,
# #     }
# 
# #     # Try digit match first, e.g., "10 slides" or "10 slide" or "10 page" or "10 pages"
# #     m_digit = re.search(r"(\d+)\s*(?:slide|page)s?", query.lower())
# #     if m_digit:
# #         slide_count = int(m_digit.group(1))
# #     else:
# #         # Try word match, e.g., "ten slides" or "ten slide"
# #         m_word = re.search(r"\b(" + "|".join(word_to_num.keys()) + r")\b\s*(?:slide|page)s?", query.lower())
# #         if m_word:
# #             slide_count = word_to_num[m_word.group(1)]
# 
# #     # Keep slide count in a safe, reasonable range (1 to 20) to avoid timeouts
# #     slide_count = max(1, min(slide_count, 20))
# 
# #     # Clean the topic string by removing the slide count specification so that the LLM only generates on the topic itself
# #     topic = query
# #     topic = re.sub(
# #         r"\b(?:with|having|of|containing)?\s*\d+\s*(?:slide|page)s?\b", "", topic, flags=re.IGNORECASE
# #     )
# #     topic = re.sub(
# #         r"\b(?:with|having|of|containing)?\s*(?:" + "|".join(word_to_num.keys()) + r")\s*(?:slide|page)s?\b",
# #         "",
# #         topic,
# #         flags=re.IGNORECASE,
# #     )
# #     topic = re.sub(r"\s+", " ", topic).strip(" .?!")
# 
# #     if not topic:
# #         topic = "Presentation"
# 
# #     logger.info(f"Generating presentation on topic: '{topic}' with slide count: {slide_count} via voice...")
# #     slides = await create_presentation_from_prompt(
# #         prompt=topic, session_id=session_id, slide_count=slide_count
# #     )
# 
# #     await bus.emit_event(
# #         "ppt_command", {"action": "reload", "slides": slides, "filename": f"AI: {topic}"}, session_id
# #     )
# 
# #     return {
# #         "status": "ok",
# #         "spoken_reply": f"I have generated a new widescreen presentation on {topic} containing {len(slides)} slides. It is now loaded on your screen.",
# #     }
# 
# 
# # async def ppt_summarize(args: dict, session_id: str) -> dict:
# #     effective_sid = resolve_sid(session_id)
# #     slides = _slide_store.get(effective_sid, []) or _slide_store.get("default", [])
# #     if not slides:
# #         return {"spoken_reply": "No presentation is loaded yet. Please upload a PowerPoint file first."}
# 
# #     lines = []
# #     for s in slides:
# #         title = s.get("title", f"Slide {s['index'] + 1}")
# #         notes = s.get("notes", "")
# #         lines.append(f"Slide {s['index'] + 1}: {title}" + (f" — {notes}" if notes else ""))
# 
# #     content = "\n".join(lines)
# #     logger.info(f"📡 [PPT TOOL] Requesting presentation summary from Ollama ({len(slides)} slides)...")
# #     summary = await _summarize_ollama(content)
# #     logger.info("📡 [PPT TOOL] Ollama summary completed successfully.")
# #     return {"spoken_reply": summary}
# 
# 
# # async def _summarize_ollama(content: str) -> str:
# #     def _call() -> str:
# #         import ollama
# 
# #         from backend.core.config import settings
# 
# #         resp = ollama.chat(
# #             model=settings.OLLAMA_MODEL,
# #             messages=[
# #                 {
# #                     "role": "system",
# #                     "content": "You are a helpful voice assistant. Summarize the presentation in 4-6 natural spoken sentences. "
# #                     "Mention the main topics and key points. No markdown, no bullet points — plain conversational speech only.",
# #                 },
# #                 {"role": "user", "content": f"Summarize this presentation:\n\n{content[:30000]}"},
# #             ],
# #             think=False,
# #             options={"num_predict": 220},
# #             stream=False,
# #         )
# #         if isinstance(resp, dict):
# #             return resp["message"]["content"].strip()
# #         return resp.message.content.strip()
# 
# #     try:
# #         return await asyncio.to_thread(_call)
# #     except Exception as e:
# #         logger.error(f"ppt_summarize ollama error: {e}")
# #         return "I wasn't able to summarize the presentation right now. Please try again."
# 
# 
# # async def ppt_qa(args: dict, session_id: str) -> dict:
# #     import re
# 
# #     query = args.get("query", "").strip()
# #     if not query:
# #         return {"spoken_reply": "I'm here! What would you like to know about the slides?", "status": "ok"}
# 
# #     effective_sid = resolve_sid(session_id)
# #     slides = _slide_store.get(effective_sid, []) or _slide_store.get("default", [])
# #     if not slides:
# #         return {
# #             "spoken_reply": "I don't have any presentation slides loaded yet. Please upload a PowerPoint file first.",
# #             "status": "ok",
# #         }
# 
# #     # Extract slide number/index
# #     # Supports word equivalents of numbers as well
# #     word_to_num = {
# #         "one": 1,
# #         "two": 2,
# #         "three": 3,
# #         "four": 4,
# #         "five": 5,
# #         "six": 6,
# #         "seven": 7,
# #         "eight": 8,
# #         "nine": 9,
# #         "ten": 10,
# #         "eleven": 11,
# #         "twelve": 12,
# #         "thirteen": 13,
# #         "fourteen": 14,
# #         "fifteen": 15,
# #         "sixteen": 16,
# #         "seventeen": 17,
# #         "eighteen": 18,
# #         "nineteen": 19,
# #         "twenty": 20,
# #     }
# 
# #     slide_index = None
# #     num_match = re.search(r"slide\s+(\d+)", query.lower())
# #     if num_match:
# #         slide_index = int(num_match.group(1)) - 1
# #     else:
# #         for word, num in word_to_num.items():
# #             if re.search(rf"\bslide\s+{word}\b", query.lower()):
# #                 slide_index = num - 1
# #                 break
# 
# #     # If no explicit slide number was matched, check if the user refers to the current/active slide
# #     if slide_index is None:
# #         current_keywords = [
# #             "current slide",
# #             "this slide",
# #             "active slide",
# #             "present slide",
# #             "current page",
# #             "this page",
# #         ]
# #         if any(kw in query.lower() for kw in current_keywords):
# #             slide_index = _current_slide.get(effective_sid, 0)
# 
# #     # If explicit or resolved slide index was matched and is valid
# #     if slide_index is not None and 0 <= slide_index < len(slides):
# #         slide = slides[slide_index]
# #         title = slide.get("title", f"Slide {slide_index + 1}")
# #         bullets_text = " ".join(slide.get("bullets", []))
# #         notes = slide.get("notes", "")
# #         img_b64 = slide.get("img_b64")
# 
# #         slide_desc = f"Slide {slide_index + 1}: {title}. Content: {bullets_text}."
# #         if notes:
# #             slide_desc += f" Speaker Notes: {notes}."
# 
# #         logger.info(f"Answering PPT Q&A for slide {slide_index + 1} using slide data.")
# 
# #         prompt = (
# #             f"The user is asking: '{query}'. This is regarding slide {slide_index + 1} "
# #             f"of the presentation, which is titled '{title}'. "
# #             f"The text content on this slide is: '{bullets_text}'. "
# #             f"Speaker notes for this slide say: '{notes}'.\n"
# #             f"Please write a friendly, clear, and comprehensive spoken response of 2-3 sentences answering the user's question directly based on this slide content."
# #         )
# 
# #         # If image is available and we have a vision-capable provider
# #         if img_b64:
# #             logger.info("Slide image is available. Calling Vision LLM.")
# #             reply = await _call_vision_llm(prompt, img_b64)
# #             if reply:
# #                 return {"spoken_reply": reply, "status": "ok", "slide_index": slide_index}
# 
# #         # Text-only fallback
# #         reply = await _call_text_llm(prompt)
# #         return {"spoken_reply": reply, "status": "ok", "slide_index": slide_index}
# 
# #     else:
# #         # General presentation Q&A
# #         all_content = []
# #         for s in slides:
# #             bullets_str = ", ".join(s.get("bullets", []))
# #             all_content.append(f"Slide {s['index'] + 1} (Title: {s.get('title', '')}): {bullets_str}")
# #         full_structure = "\n".join(all_content)
# 
# #         prompt = (
# #             f"The user is asking a question about the PowerPoint presentation: '{query}'. "
# #             f"Here is the text structure extracted from the presentation:\n"
# #             f"{full_structure[:30000]}\n"
# #             f"Please write a friendly, concise spoken response of 2-3 sentences answering their question directly based on the presentation slides."
# #         )
# 
# #         logger.info("General PPT Q&A. Answering using full presentation text structure.")
# #         reply = await _call_text_llm(prompt)
# #         return {"spoken_reply": reply, "status": "ok"}
# 
# 
# # async def _call_vision_llm(prompt: str, img_b64: str) -> str | None:
# #     from backend.core.config import settings
# 
# #     # Try Gemini Vision first (always preferred)
# #     if settings.GEMINI_API_KEY:
# #         try:
# #             import base64
# 
# #             import google.generativeai as genai
# 
# #             genai.configure(api_key=settings.GEMINI_API_KEY)
# #             # Use gemini-2.5-flash as it is the absolute latest and highest-capability Gemini model for vision and MCP Q&A
# #             model = genai.GenerativeModel("gemini-2.5-flash")
# #             img_bytes = base64.b64decode(img_b64)
# #             resp = await model.generate_content_async([prompt, {"mime_type": "image/png", "data": img_bytes}])
# #             if resp.text:
# #                 return resp.text.strip()
# #         except Exception as e:
# #             logger.warning(f"Gemini vision PPT Q&A failed: {e}")
# 
# #     # Try Groq vision if key is active
# #     if settings.GROQ_API_KEY:
# #         try:
# #             from groq import AsyncGroq
# 
# #             client = AsyncGroq(api_key=settings.GROQ_API_KEY)
# #             resp = await client.chat.completions.create(
# #                 model="llama-3.2-11b-vision-preview",
# #                 messages=[
# #                     {
# #                         "role": "user",
# #                         "content": [
# #                             {"type": "text", "text": prompt},
# #                             {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{img_b64}"}},
# #                         ],
# #                     }
# #                 ],
# #                 max_tokens=200,
# #             )
# #             if resp.choices[0].message.content:
# #                 return resp.choices[0].message.content.strip()
# #         except Exception as e:
# #             logger.warning(f"Groq vision PPT Q&A failed: {e}")
# 
# #     return None
# 
# 
# # async def _call_text_llm(prompt: str) -> str:
# #     from backend.core.config import settings
# 
# #     # Try Groq first
# #     if settings.GROQ_API_KEY:
# #         try:
# #             from groq import AsyncGroq
# 
# #             client = AsyncGroq(api_key=settings.GROQ_API_KEY)
# #             resp = await client.chat.completions.create(
# #                 model="llama-3.1-8b-instant",
# #                 messages=[
# #                     {
# #                         "role": "system",
# #                         "content": "You are a helpful voice assistant. Answer concisely in 1-3 spoken sentences.",
# #                     },
# #                     {"role": "user", "content": prompt},
# #                 ],
# #                 max_tokens=200,
# #             )
# #             if resp.choices[0].message.content:
# #                 return resp.choices[0].message.content.strip()
# #         except Exception as e:
# #             logger.warning(f"Groq PPT text Q&A failed: {e}")
# 
# #     # Try Gemini
# #     if settings.GEMINI_API_KEY:
# #         try:
# #             import google.generativeai as genai
# 
# #             genai.configure(api_key=settings.GEMINI_API_KEY)
# #             # Use gemini-2.5-flash as it is the absolute latest and highest-capability Gemini model for text Q&A
# #             model = genai.GenerativeModel("gemini-2.5-flash")
# #             resp = await model.generate_content_async(prompt)
# #             if resp.text:
# #                 return resp.text.strip()
# #         except Exception as e:
# #             logger.warning(f"Gemini PPT text Q&A failed: {e}")
# 
# #     # Fallback to local Ollama
# #     try:
# #         import httpx
# 
# #         url = f"{settings.OLLAMA_BASE_URL.rstrip('/')}/api/generate"
# #         async with httpx.AsyncClient(timeout=25.0) as client:
# #             resp = await client.post(
# #                 url,
# #                 json={
# #                     "model": settings.OLLAMA_MODEL,
# #                     "prompt": f"System: You are a helpful voice assistant. Answer concisely in 1-3 spoken sentences.\n\nUser: {prompt}",
# #                     "stream": False,
# #                 },
# #             )
# #             if resp.status_code == 200:
# #                 return resp.json().get("response", "").strip()
# #     except Exception as e:
# #         logger.error(f"Ollama local PPT text Q&A failed: {e}")
# 
# #     return "I analyzed the presentation slide but was unable to generate a response. Please check your AI API keys."
# 
# 
# # async def _generate_single_slide_content(topic: str) -> dict:
# #     import json
# #     import re
# 
# #     import httpx
# 
# #     from backend.core.config import settings
# 
# #     system_content = (
# #         "You are an expert presentation designer. Create a SINGLE highly professional slide based on the user's topic.\n"
# #         "You MUST respond ONLY with a valid JSON object. Do not write any markdown, "
# #         "explanations, or backticks outside the JSON. The JSON must follow this exact schema:\n"
# #         "{\n"
# #         '  "title": "Slide Title",\n'
# #         '  "bullets": ["Detailed, informative bullet point 1", "Detailed, informative bullet point 2"],\n'
# #         '  "notes": "Comprehensive speaker notes detailing the slide concepts"\n'
# #         "}\n"
# #         "Keep bullets concise and professional (under 12 words each, max 4 bullets per slide)."
# #     )
# 
# #     parsed = None
# 
# #     # 1. Try Gemini
# #     if settings.GEMINI_API_KEY:
# #         try:
# #             import google.generativeai as genai
# 
# #             genai.configure(api_key=settings.GEMINI_API_KEY)
# #             model = genai.GenerativeModel("gemini-2.5-flash")
# #             response = await model.generate_content_async(
# #                 f"Create a single professional slide on the topic: '{topic}'",
# #                 generation_config={
# #                     "response_mime_type": "application/json",
# #                     "system_instruction": system_content,
# #                 },
# #             )
# #             parsed = json.loads(response.text.strip())
# #         except Exception as e:
# #             logger.warning(f"Cloud Gemini single slide generation failed: {e}")
# 
# #     # 2. Try Groq
# #     if not parsed and settings.GROQ_API_KEY:
# #         try:
# #             from groq import AsyncGroq
# 
# #             client = AsyncGroq(api_key=settings.GROQ_API_KEY)
# #             resp = await client.chat.completions.create(
# #                 model="llama-3.3-70b-versatile",
# #                 messages=[
# #                     {"role": "system", "content": system_content},
# #                     {"role": "user", "content": f"Create a single slide on: '{topic}'"},
# #                 ],
# #                 response_format={"type": "json_object"},
# #             )
# #             parsed = json.loads(resp.choices[0].message.content)
# #         except Exception as e:
# #             logger.warning(f"Cloud Groq single slide generation failed: {e}")
# 
# #     # 3. Fallback to local Ollama
# #     if not parsed:
# #         try:
# #             url = f"{settings.OLLAMA_BASE_URL.rstrip('/')}/api/chat"
# #             payload = {
# #                 "model": settings.OLLAMA_MODEL,
# #                 "stream": False,
# #                 "messages": [
# #                     {"role": "system", "content": system_content},
# #                     {"role": "user", "content": f"Create a single slide on: '{topic}'"},
# #                 ],
# #                 "options": {"temperature": 0.3, "num_predict": 500},
# #                 "format": "json",
# #             }
# #             async with httpx.AsyncClient(timeout=40.0) as client:
# #                 resp = await client.post(url, json=payload)
# #                 if resp.status_code == 200:
# #                     result_data = resp.json()
# #                     raw_content = result_data.get("message", {}).get("content", "").strip()
# #                     try:
# #                         parsed = json.loads(raw_content)
# #                     except Exception:
# #                         json_match = re.search(r"(\{.*\})", raw_content, re.DOTALL)
# #                         if json_match:
# #                             parsed = json.loads(json_match.group(1))
# #         except Exception as e:
# #             logger.error(f"Local Ollama single slide generation failed: {e}")
# 
# #     # Normalize output
# #     if parsed and isinstance(parsed, dict):
# #         title = parsed.get("title", topic.title())
# #         bullets = parsed.get("bullets", [])
# #         if isinstance(bullets, str):
# #             bullets = [bullets]
# #         elif not isinstance(bullets, list):
# #             bullets = []
# #         notes = parsed.get("notes", "")
# #         return {"title": str(title), "bullets": [str(b) for b in bullets], "notes": str(notes)}
# 
# #     # Fallback template slide
# #     return {
# #         "title": topic.title(),
# #         "bullets": [
# #             f"Overview of {topic}",
# #             "Key definitions and background concepts",
# #             "Primary challenges and modern implications",
# #         ],
# #         "notes": f"This slide presents key details about {topic}.",
# #     }
# 
# 
# # async def ppt_add_slide(args: dict, session_id: str) -> dict:
# #     from backend.api.ppt import save_and_render_pptx
# #     from backend.core.slide_store import resolve_sid
# #     from backend.queues.bus import bus
# 
# #     effective_sid = resolve_sid(session_id)
# #     slides = list(_slide_store.get(effective_sid, []))
# 
# #     topic = args.get("topic", "").strip(" .?!")
# #     if not topic:
# #         topic = "New Slide"
# 
# #     logger.info(f"Adding a slide on topic '{topic}' via voice...")
# 
# #     # Generate content
# #     if topic != "New Slide":
# #         new_slide = await _generate_single_slide_content(topic)
# #     else:
# #         new_slide = {
# #             "title": "New Slide",
# #             "bullets": ["Write content or use voice to edit this slide"],
# #             "notes": "",
# #         }
# 
# #     # Recompile plain slide data
# #     slides_to_save = []
# #     for s in slides:
# #         slides_to_save.append(
# #             {"title": s.get("title", ""), "bullets": s.get("bullets", []), "notes": s.get("notes", "")}
# #         )
# #     slides_to_save.append(new_slide)
# 
# #     # Save and render
# #     rendered = save_and_render_pptx(slides_to_save, effective_sid, topic=topic)
# #     new_idx = len(rendered) - 1
# #     _current_slide[effective_sid] = new_idx
# 
# #     await bus.emit_event(
# #         "ppt_command",
# #         {
# #             "action": "reload",
# #             "slides": rendered,
# #             "filename": f"Presentation ({len(rendered)} slides)",
# #             "index": new_idx,
# #             "preserveCurrent": False,
# #         },
# #         session_id,
# #     )
# 
# #     return {
# #         "status": "ok",
# #         "spoken_reply": f"I've added a slide about {new_slide['title']} and navigated to it.",
# #     }
# 
# 
# # async def ppt_edit_slide(args: dict, session_id: str) -> dict:
# #     from backend.api.ppt import save_and_render_pptx
# #     from backend.core.slide_store import resolve_sid
# #     from backend.queues.bus import bus
# 
# #     effective_sid = resolve_sid(session_id)
# #     slides = list(_slide_store.get(effective_sid, []))
# #     if not slides:
# #         return {"status": "error", "spoken_reply": "There is no active presentation to edit."}
# 
# #     current_idx = _current_slide.get(effective_sid, 0)
# #     if current_idx < 0 or current_idx >= len(slides):
# #         current_idx = 0
# 
# #     title = args.get("title")
# #     bullets = args.get("bullets")
# #     add_bullet = args.get("add_bullet")
# #     notes = args.get("notes")
# 
# #     spoken_reply = "Updated the slide successfully."
# 
# #     if title is not None:
# #         slides[current_idx]["title"] = title
# #         spoken_reply = f"I have updated the slide title to '{title}'."
# #     if bullets is not None:
# #         slides[current_idx]["bullets"] = bullets
# #         spoken_reply = "I have updated the bullet points on this slide."
# #     elif add_bullet is not None:
# #         if "bullets" not in slides[current_idx] or not isinstance(slides[current_idx]["bullets"], list):
# #             slides[current_idx]["bullets"] = []
# #         slides[current_idx]["bullets"].append(add_bullet)
# #         spoken_reply = f"I've added a bullet point about '{add_bullet}' to this slide."
# #     if notes is not None:
# #         slides[current_idx]["notes"] = notes
# #         spoken_reply = "I have updated the speaker notes for this slide."
# 
# #     # Recompile plain slide data
# #     slides_to_save = []
# #     for s in slides:
# #         slides_to_save.append(
# #             {"title": s.get("title", ""), "bullets": s.get("bullets", []), "notes": s.get("notes", "")}
# #         )
# 
# #     rendered = save_and_render_pptx(slides_to_save, effective_sid, topic="Presentation")
# #     _current_slide[effective_sid] = current_idx
# 
# #     await bus.emit_event(
# #         "ppt_command",
# #         {
# #             "action": "reload",
# #             "slides": rendered,
# #             "filename": f"Presentation ({len(rendered)} slides)",
# #             "index": current_idx,
# #             "preserveCurrent": True,
# #         },
# #         session_id,
# #     )
# 
# #     return {"status": "ok", "spoken_reply": spoken_reply}
# 
# 
# # async def ppt_clear_presentation(args: dict, session_id: str) -> dict:
# #     from backend.api.ppt import save_and_render_pptx
# #     from backend.core.slide_store import resolve_sid
# #     from backend.queues.bus import bus
# 
# #     effective_sid = resolve_sid(session_id)
# #     default_slides = [
# #         {
# #             "title": "New Presentation",
# #             "bullets": ["Voice-Driven Interactive Presentation", "Start speaking or clicking to add slides"],
# #             "notes": "Welcome to your new interactive presentation deck.",
# #         }
# #     ]
# 
# #     rendered = save_and_render_pptx(default_slides, effective_sid, topic="New Presentation")
# #     _current_slide[effective_sid] = 0
# 
# #     await bus.emit_event(
# #         "ppt_command",
# #         {
# #             "action": "reload",
# #             "slides": rendered,
# #             "filename": "New Presentation",
# #             "index": 0,
# #             "preserveCurrent": False,
# #         },
# #         session_id,
# #     )
# 
# #     return {
# #         "status": "ok",
# #         "spoken_reply": "I have started a new presentation deck with a blank title slide. What topic would you like to add a slide about?",
# #     }
# 
# 
# # async def ppt_improvise_slide(args: dict, session_id: str) -> dict:
# #     from backend.api.ppt import improvise_slide_content, save_and_render_pptx
# #     from backend.core.slide_store import resolve_sid
# #     from backend.queues.bus import bus
# 
# #     effective_sid = resolve_sid(session_id)
# #     slides = list(_slide_store.get(effective_sid, []))
# #     if not slides:
# #         return {"status": "error", "spoken_reply": "There is no active presentation to improvise."}
# 
# #     current_idx = _current_slide.get(effective_sid, 0)
# #     if current_idx < 0 or current_idx >= len(slides):
# #         current_idx = 0
# 
# #     prompt = args.get("prompt", "").strip(" .?!")
# #     if not prompt:
# #         return {"status": "error", "spoken_reply": "What would you like me to improve on this slide?"}
# 
# #     logger.info(f"Improvising slide {current_idx + 1} with prompt: '{prompt}' via voice...")
# 
# #     try:
# #         current_slide = slides[current_idx]
# #         improved = await improvise_slide_content(current_slide, prompt)
# 
# #         # Update text fields but preserve any images!
# #         slides[current_idx]["title"] = improved["title"]
# #         slides[current_idx]["bullets"] = improved["bullets"]
# #         slides[current_idx]["notes"] = improved["notes"]
# 
# #         # Recompile plain slide data
# #         slides_to_save = []
# #         for s in slides:
# #             slides_to_save.append(
# #                 {
# #                     "title": s.get("title", ""),
# #                     "bullets": s.get("bullets", []),
# #                     "notes": s.get("notes", ""),
# #                     "images": s.get("images", []),
# #                 }
# #             )
# 
# #         rendered = save_and_render_pptx(slides_to_save, effective_sid, topic="Presentation")
# #         _current_slide[effective_sid] = current_idx
# 
# #         await bus.emit_event(
# #             "ppt_command",
# #             {
# #                 "action": "reload",
# #                 "slides": rendered,
# #                 "filename": f"Presentation ({len(rendered)} slides)",
# #                 "index": current_idx,
# #                 "preserveCurrent": True,
# #             },
# #             session_id,
# #         )
# 
# #         return {
# #             "status": "ok",
# #             "spoken_reply": f"I've improvised this slide according to your request: {prompt}.",
# #         }
# #     except Exception as e:
# #         logger.error(f"Error improvising slide via voice: {e}")
# #         return {
# #             "status": "error",
# #             "spoken_reply": "I encountered an error trying to improvise the slide. Please try again.",
# #         }
# 
# 
# import asyncio
# import logging
# 
# from backend.core.slide_store import _current_slide, _slide_store, resolve_sid
# from backend.queues.bus import bus
# 
# """PPT tools — navigate + jump to slide by number or title + summarize."""
# 
# logger = logging.getLogger("pilot.tools.ppt")
# 
# 
# 
# 
# # ---- superseded by GD-template implementation below (migrated from
# # ---- syugesh/pilot-voice-agent-backend @ 2d3df73) — kept for reference ----
# # async def ppt_navigate(args: dict, session_id: str) -> dict:
# #     direction = args.get("direction","next")
# #     effective_sid = resolve_sid(session_id)
# #     slides = _slide_store.get(effective_sid, [])
# #     total = len(slides)
# #     current = _current_slide.get(effective_sid, 0)
# # 
# #     if direction == "next":
# #         if total > 0 and current >= total - 1:
# #             return {
# #                 "spoken_reply": f"You've reached the last slide — slide {total} of {total}. That's the end of the presentation."
# #             }
# #         _current_slide[effective_sid] = min(current + 1, max(total - 1, 0))
# #     elif direction == "prev":
# #         if current <= 0:
# #             return {"spoken_reply": "You're already on the first slide."}
# #         _current_slide[effective_sid] = max(current - 1, 0)
# #     elif direction == "first":
# #         _current_slide[effective_sid] = 0
# #     elif direction == "last":
# #         _current_slide[effective_sid] = max(total - 1, 0)
# # 
# #     await bus.emit_event("ppt_command", {"action": direction}, session_id)
# #     return {"status": "ok", "direction": direction, "index": _current_slide.get(effective_sid, 0)}
# # 
# # 
# # async def ppt_jump_to_title(args: dict, session_id: str) -> dict:
# #     query = args.get("query", "")
# #     slide_number = args.get("slide_number")  # already 0-indexed if from keyword fallback
# #     import re
# # 
# #     from backend.queues.bus import bus
# # 
# #     effective_sid = resolve_sid(session_id)
# #     slides = _slide_store.get(effective_sid, [])
# # 
# #     # If explicit slide number given, use directly
# #     if slide_number is not None:
# #         idx = int(slide_number)
# #         _current_slide[effective_sid] = idx
# #         await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
# #         return {"status": "ok", "index": idx, "title": f"Slide {idx + 1}"}
# # 
# #     q = query.lower()
# # 
# #     # Numeric match in query — only allow within actual deck bounds
# #     m = re.search(r"\b(\d+)\b", q)
# #     if m:
# #         idx = int(m.group(1)) - 1
# #         if 0 <= idx < len(slides):
# #             _current_slide[effective_sid] = idx
# #             await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
# #             return {"status": "ok", "index": idx}
# # 
# #     # Title fuzzy match
# #     best_idx, best_score = 0, 0
# #     for s in slides:
# #         score = sum(1 for w in q.split() if w in s.get("title", "").lower())
# #         if score > best_score:
# #             best_score = score
# #             best_idx = s["index"]
# # 
# #     _current_slide[effective_sid] = best_idx
# #     await bus.emit_event("ppt_command", {"action": "goto", "index": best_idx}, session_id)
# #     return {"status": "ok", "index": best_idx}
# # 
# # 
# 
# def _prepare_slides_for_save(slides: list) -> list[dict]:
#     return [
#         {
#             "title": s.get("title", ""),
#             "bullets": s.get("bullets", []),
#             "notes": s.get("notes", ""),
#             "images": s.get("images", []),
#             "font_name": s.get("font_name", "Arial"),
#             "title_bold": s.get("title_bold", True),
#             "title_italic": s.get("title_italic", False),
#             "bullet_bold": s.get("bullet_bold", False),
#             "bullet_italic": s.get("bullet_italic", False),
#         }
#         for s in slides
#     ]
# 
# 
# 
# # ---- superseded by GD-template implementation below ----
# # async def ppt_delete_slide(args: dict, session_id: str) -> dict:
# #     from backend.api.ppt import _current_slide, _slide_store, resolve_sid, save_and_render_pptx
# #     from backend.queues.bus import bus
# # 
# #     effective_sid = resolve_sid(session_id)
# #     slides = list(_slide_store.get(effective_sid, []))
# #     if not slides:
# #         return {"status": "error", "spoken_reply": "There are no slides in the presentation to delete."}
# # 
# #     current_idx = _current_slide.get(effective_sid, 0)
# #     if current_idx < 0 or current_idx >= len(slides):
# #         current_idx = 0
# # 
# #     slides.pop(current_idx)
# # 
# #     # Re-build slide list preserving images and style details
# #     slides_to_save = _prepare_slides_for_save(slides)
# # 
# #     try:
# #         rendered = save_and_render_pptx(slides_to_save, effective_sid, topic="Presentation")
# #         new_idx = min(current_idx, max(len(rendered) - 1, 0))
# #         _current_slide[effective_sid] = new_idx
# # 
# #         # Emit reload event
# #         await bus.emit_event(
# #             "ppt_command",
# #             {
# #                 "action": "reload",
# #                 "slides": rendered,
# #                 "filename": f"Presentation ({len(rendered)} slides)",
# #                 "index": new_idx,
# #                 "preserveCurrent": False,
# #             },
# #             session_id,
# #         )
# # 
# #         return {"status": "ok", "spoken_reply": "I have deleted the current slide."}
# #     except Exception as e:
# #         return {"status": "error", "spoken_reply": f"Failed to delete slide: {e}"}
# # 
# # 
# 
# async def ppt_create_slides(args: dict, session_id: str) -> dict:
#     import re
# 
#     from backend.api.ppt import create_presentation_from_prompt
#     from backend.queues.bus import bus
# 
#     query = args.get("query", "presentation")
# 
#     # ── DYNAMIC SLIDE COUNT PARSING ──
#     # Default is 5 slides
#     slide_count = 5
# 
#     word_to_num = {
#         "one": 1,
#         "two": 2,
#         "three": 3,
#         "four": 4,
#         "five": 5,
#         "six": 6,
#         "seven": 7,
#         "eight": 8,
#         "nine": 9,
#         "ten": 10,
#         "eleven": 11,
#         "twelve": 12,
#         "thirteen": 13,
#         "fourteen": 14,
#         "fifteen": 15,
#         "sixteen": 16,
#         "seventeen": 17,
#         "eighteen": 18,
#         "nineteen": 19,
#         "twenty": 20,
#     }
# 
#     # Try digit match first, e.g., "10 slides" or "10 slide" or "10 page" or "10 pages"
#     m_digit = re.search(r"(\d+)\s*(?:slide|page)s?", query.lower())
#     if m_digit:
#         slide_count = int(m_digit.group(1))
#     else:
#         # Try word match, e.g., "ten slides" or "ten slide"
#         m_word = re.search(r"\b(" + "|".join(word_to_num.keys()) + r")\b\s*(?:slide|page)s?", query.lower())
#         if m_word:
#             slide_count = word_to_num[m_word.group(1)]
# 
#     # Keep slide count in a safe, reasonable range (1 to 20) to avoid timeouts
#     slide_count = max(1, min(slide_count, 20))
# 
#     # Clean the topic string by removing the slide count specification so that the LLM only generates on the topic itself
#     topic = query
#     topic = re.sub(
#         r"\b(?:with|having|of|containing)?\s*\d+\s*(?:slide|page)s?\b", "", topic, flags=re.IGNORECASE
#     )
#     topic = re.sub(
#         r"\b(?:with|having|of|containing)?\s*(?:" + "|".join(word_to_num.keys()) + r")\s*(?:slide|page)s?\b",
#         "",
#         topic,
#         flags=re.IGNORECASE,
#     )
#     topic = re.sub(r"\s+", " ", topic).strip(" .?!")
# 
#     if not topic:
#         topic = "Presentation"
# 
#     logger.info(f"Generating presentation on topic: '{topic}' with slide count: {slide_count} via voice...")
#     slides = await create_presentation_from_prompt(
#         prompt=topic, session_id=session_id, slide_count=slide_count
#     )
# 
#     await bus.emit_event(
#         "ppt_command", {"action": "reload", "slides": slides, "filename": f"AI: {topic}"}, session_id
#     )
# 
#     return {
#         "status": "ok",
#         "spoken_reply": f"I have generated a new widescreen presentation on {topic} containing {len(slides)} slides. It is now loaded on your screen.",
#     }
# 
# 
# 
# # ---- superseded by GD-template implementation below ----
# # async def ppt_summarize(args: dict, session_id: str) -> dict:
# #     effective_sid = resolve_sid(session_id)
# #     slides = _slide_store.get(effective_sid, []) or _slide_store.get("default", [])
# #     if not slides:
# #         return {"spoken_reply": "No presentation is loaded yet. Please upload a PowerPoint file first."}
# # 
# #     lines = []
# #     for s in slides:
# #         title = s.get("title", f"Slide {s['index'] + 1}")
# #         notes = s.get("notes", "")
# #         lines.append(f"Slide {s['index'] + 1}: {title}" + (f" — {notes}" if notes else ""))
# # 
# #     content = "\n".join(lines)
# #     logger.info(f"📡 [PPT TOOL] Requesting presentation summary from Ollama ({len(slides)} slides)...")
# #     summary = await _summarize_ollama(content)
# #     logger.info("📡 [PPT TOOL] Ollama summary completed successfully.")
# #     return {"spoken_reply": summary}
# # 
# # 
# # async def _summarize_ollama(content: str) -> str:
# #     def _call() -> str:
# #         import ollama
# # 
# #         from backend.core.config import settings
# # 
# #         resp = ollama.chat(
# #             model=settings.OLLAMA_MODEL,
# #             messages=[
# #                 {
# #                     "role": "system",
# #                     "content": "You are a helpful voice assistant. Summarize the presentation in 4-6 natural spoken sentences. "
# #                     "Mention the main topics and key points. No markdown, no bullet points — plain conversational speech only.",
# #                 },
# #                 {"role": "user", "content": f"Summarize this presentation:\n\n{content[:30000]}"},
# #             ],
# #             think=False,
# #             options={"num_predict": 220},
# #             stream=False,
# #         )
# #         if isinstance(resp, dict):
# #             return resp["message"]["content"].strip()
# #         return resp.message.content.strip()
# # 
# #     try:
# #         return await asyncio.to_thread(_call)
# #     except Exception as e:
# #         logger.error(f"ppt_summarize ollama error: {e}")
# #         return "I wasn't able to summarize the presentation right now. Please try again."
# # 
# # 
# 
# async def ppt_qa(args: dict, session_id: str) -> dict:
#     import re
# 
#     query = args.get("query", "").strip()
#     if not query:
#         return {"spoken_reply": "I'm here! What would you like to know about the slides?", "status": "ok"}
# 
#     effective_sid = resolve_sid(session_id)
#     slides = _slide_store.get(effective_sid, []) or _slide_store.get("default", [])
#     if not slides:
#         return {
#             "spoken_reply": "I don't have any presentation slides loaded yet. Please upload a PowerPoint file first.",
#             "status": "ok",
#         }
# 
#     # Extract slide number/index
#     # Supports word equivalents of numbers as well
#     word_to_num = {
#         "one": 1,
#         "two": 2,
#         "three": 3,
#         "four": 4,
#         "five": 5,
#         "six": 6,
#         "seven": 7,
#         "eight": 8,
#         "nine": 9,
#         "ten": 10,
#         "eleven": 11,
#         "twelve": 12,
#         "thirteen": 13,
#         "fourteen": 14,
#         "fifteen": 15,
#         "sixteen": 16,
#         "seventeen": 17,
#         "eighteen": 18,
#         "nineteen": 19,
#         "twenty": 20,
#     }
# 
#     slide_index = None
#     num_match = re.search(r"slide\s+(\d+)", query.lower())
#     if num_match:
#         slide_index = int(num_match.group(1)) - 1
#     else:
#         for word, num in word_to_num.items():
#             if re.search(rf"\bslide\s+{word}\b", query.lower()):
#                 slide_index = num - 1
#                 break
# 
#     # If no explicit slide number was matched, check if the user refers to the current/active slide
#     if slide_index is None:
#         current_keywords = [
#             "current slide",
#             "this slide",
#             "active slide",
#             "present slide",
#             "current page",
#             "this page",
#         ]
#         if any(kw in query.lower() for kw in current_keywords):
#             slide_index = _current_slide.get(effective_sid, 0)
# 
#     # If explicit or resolved slide index was matched and is valid
#     if slide_index is not None and 0 <= slide_index < len(slides):
#         slide = slides[slide_index]
#         title = slide.get("title", f"Slide {slide_index + 1}")
#         bullets_text = " ".join(slide.get("bullets", []))
#         notes = slide.get("notes", "")
#         img_b64 = slide.get("img_b64")
# 
#         slide_desc = f"Slide {slide_index + 1}: {title}. Content: {bullets_text}."
#         if notes:
#             slide_desc += f" Speaker Notes: {notes}."
# 
#         logger.info(f"Answering PPT Q&A for slide {slide_index + 1} using slide data.")
# 
#         prompt = (
#             f"The user is asking: '{query}'. This is regarding slide {slide_index + 1} "
#             f"of the presentation, which is titled '{title}'. "
#             f"The text content on this slide is: '{bullets_text}'. "
#             f"Speaker notes for this slide say: '{notes}'.\n"
#             f"Please write a friendly, clear, and comprehensive spoken response of 2-3 sentences answering the user's question directly based on this slide content."
#         )
# 
#         # If image is available and we have a vision-capable provider
#         if img_b64:
#             logger.info("Slide image is available. Calling Vision LLM.")
#             reply = await _call_vision_llm(prompt, img_b64)
#             if reply:
#                 return {"spoken_reply": reply, "status": "ok", "slide_index": slide_index}
# 
#         # Text-only fallback
#         reply = await _call_text_llm(prompt)
#         return {"spoken_reply": reply, "status": "ok", "slide_index": slide_index}
# 
#     else:
#         # General presentation Q&A
#         all_content = []
#         for s in slides:
#             bullets_str = ", ".join(s.get("bullets", []))
#             all_content.append(f"Slide {s['index'] + 1} (Title: {s.get('title', '')}): {bullets_str}")
#         full_structure = "\n".join(all_content)
# 
#         prompt = (
#             f"The user is asking a question about the PowerPoint presentation: '{query}'. "
#             f"Here is the text structure extracted from the presentation:\n"
#             f"{full_structure[:30000]}\n"
#             f"Please write a friendly, concise spoken response of 2-3 sentences answering their question directly based on the presentation slides."
#         )
# 
#         logger.info("General PPT Q&A. Answering using full presentation text structure.")
#         reply = await _call_text_llm(prompt)
#         return {"spoken_reply": reply, "status": "ok"}
# 
# 
# async def _call_vision_llm(prompt: str, img_b64: str) -> str | None:
#     from backend.core.config import settings
# 
#     # Try Gemini Vision first (always preferred)
#     if settings.GEMINI_API_KEY:
#         try:
#             import base64
# 
#             import google.generativeai as genai
# 
#             genai.configure(api_key=settings.GEMINI_API_KEY)
#             # Use gemini-2.5-flash as it is the absolute latest and highest-capability Gemini model for vision and MCP Q&A
#             model = genai.GenerativeModel("gemini-2.5-flash")
#             img_bytes = base64.b64decode(img_b64)
#             resp = await model.generate_content_async([prompt, {"mime_type": "image/png", "data": img_bytes}])
#             if resp.text:
#                 return resp.text.strip()
#         except Exception as e:
#             logger.warning(f"Gemini vision PPT Q&A failed: {e}")
# 
#     # Try Groq vision if key is active
#     if settings.GROQ_API_KEY:
#         try:
#             from groq import AsyncGroq
# 
#             client = AsyncGroq(api_key=settings.GROQ_API_KEY)
#             resp = await client.chat.completions.create(
#                 model="llama-3.2-11b-vision-preview",
#                 messages=[
#                     {
#                         "role": "user",
#                         "content": [
#                             {"type": "text", "text": prompt},
#                             {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{img_b64}"}},
#                         ],
#                     }
#                 ],
#                 max_tokens=200,
#             )
#             if resp.choices[0].message.content:
#                 return resp.choices[0].message.content.strip()
#         except Exception as e:
#             logger.warning(f"Groq vision PPT Q&A failed: {e}")
# 
#     return None
# 
# 
# async def _call_text_llm(prompt: str) -> str:
#     from backend.core.config import settings
# 
#     # Try Groq first
#     if settings.GROQ_API_KEY:
#         try:
#             from groq import AsyncGroq
# 
#             client = AsyncGroq(api_key=settings.GROQ_API_KEY)
#             resp = await client.chat.completions.create(
#                 model="llama-3.1-8b-instant",
#                 messages=[
#                     {
#                         "role": "system",
#                         "content": "You are a helpful voice assistant. Answer concisely in 1-3 spoken sentences.",
#                     },
#                     {"role": "user", "content": prompt},
#                 ],
#                 max_tokens=200,
#             )
#             if resp.choices[0].message.content:
#                 return resp.choices[0].message.content.strip()
#         except Exception as e:
#             logger.warning(f"Groq PPT text Q&A failed: {e}")
# 
#     # Try Gemini
#     if settings.GEMINI_API_KEY:
#         try:
#             import google.generativeai as genai
# 
#             genai.configure(api_key=settings.GEMINI_API_KEY)
#             # Use gemini-2.5-flash as it is the absolute latest and highest-capability Gemini model for text Q&A
#             model = genai.GenerativeModel("gemini-2.5-flash")
#             resp = await model.generate_content_async(prompt)
#             if resp.text:
#                 return resp.text.strip()
#         except Exception as e:
#             logger.warning(f"Gemini PPT text Q&A failed: {e}")
# 
#     # Fallback to local Ollama
#     try:
#         import httpx
# 
#         url = f"{settings.OLLAMA_BASE_URL.rstrip('/')}/api/generate"
#         async with httpx.AsyncClient(timeout=25.0) as client:
#             resp = await client.post(
#                 url,
#                 json={
#                     "model": settings.OLLAMA_MODEL,
#                     "prompt": f"System: You are a helpful voice assistant. Answer concisely in 1-3 spoken sentences.\n\nUser: {prompt}",
#                     "stream": False,
#                 },
#             )
#             if resp.status_code == 200:
#                 return resp.json().get("response", "").strip()
#     except Exception as e:
#         logger.error(f"Ollama local PPT text Q&A failed: {e}")
# 
#     return "I analyzed the presentation slide but was unable to generate a response. Please check your AI API keys."
# 
# 
# 
# # ---- superseded by GD-template implementation below ----
# # async def _generate_single_slide_content(topic: str) -> dict:
# #     import json
# #     import re
# # 
# #     import httpx
# # 
# #     from backend.core.config import settings
# # 
# #     system_content = (
# #         "You are an expert presentation designer. Create a SINGLE highly professional slide based on the user's topic.\n"
# #         "You MUST respond ONLY with a valid JSON object. Do not write any markdown, "
# #         "explanations, or backticks outside the JSON. The JSON must follow this exact schema:\n"
# #         "{\n"
# #         '  "title": "Slide Title",\n'
# #         '  "bullets": ["Detailed, informative bullet point 1", "Detailed, informative bullet point 2"],\n'
# #         '  "notes": "Comprehensive speaker notes detailing the slide concepts"\n'
# #         "}\n"
# #         "Keep bullets concise and professional (under 12 words each, max 4 bullets per slide)."
# #     )
# # 
# #     parsed = None
# # 
# #     # 1. Try Gemini
# #     if settings.GEMINI_API_KEY:
# #         try:
# #             import google.generativeai as genai
# # 
# #             genai.configure(api_key=settings.GEMINI_API_KEY)
# #             model = genai.GenerativeModel("gemini-2.5-flash")
# #             response = await model.generate_content_async(
# #                 f"Create a single professional slide on the topic: '{topic}'",
# #                 generation_config={
# #                     "response_mime_type": "application/json",
# #                     "system_instruction": system_content,
# #                 },
# #             )
# #             parsed = json.loads(response.text.strip())
# #         except Exception as e:
# #             logger.warning(f"Cloud Gemini single slide generation failed: {e}")
# # 
# #     # 2. Try Groq
# #     if not parsed and settings.GROQ_API_KEY:
# #         try:
# #             from groq import AsyncGroq
# # 
# #             client = AsyncGroq(api_key=settings.GROQ_API_KEY)
# #             resp = await client.chat.completions.create(
# #                 model="llama-3.3-70b-versatile",
# #                 messages=[
# #                     {"role": "system", "content": system_content},
# #                     {"role": "user", "content": f"Create a single slide on: '{topic}'"},
# #                 ],
# #                 response_format={"type": "json_object"},
# #             )
# #             parsed = json.loads(resp.choices[0].message.content)
# #         except Exception as e:
# #             logger.warning(f"Cloud Groq single slide generation failed: {e}")
# # 
# #     # 3. Fallback to local Ollama
# #     if not parsed:
# #         try:
# #             url = f"{settings.OLLAMA_BASE_URL.rstrip('/')}/api/chat"
# #             payload = {
# #                 "model": settings.OLLAMA_MODEL,
# #                 "stream": False,
# #                 "messages": [
# #                     {"role": "system", "content": system_content},
# #                     {"role": "user", "content": f"Create a single slide on: '{topic}'"},
# #                 ],
# #                 "options": {"temperature": 0.3, "num_predict": 500},
# #                 "format": "json",
# #             }
# #             async with httpx.AsyncClient(timeout=40.0) as client:
# #                 resp = await client.post(url, json=payload)
# #                 if resp.status_code == 200:
# #                     result_data = resp.json()
# #                     raw_content = result_data.get("message", {}).get("content", "").strip()
# #                     try:
# #                         parsed = json.loads(raw_content)
# #                     except Exception:
# #                         json_match = re.search(r"(\{.*\})", raw_content, re.DOTALL)
# #                         if json_match:
# #                             parsed = json.loads(json_match.group(1))
# #         except Exception as e:
# #             logger.error(f"Local Ollama single slide generation failed: {e}")
# # 
# #     # Normalize output
# #     if parsed and isinstance(parsed, dict):
# #         title = parsed.get("title", topic.title())
# #         bullets = parsed.get("bullets", [])
# #         if isinstance(bullets, str):
# #             bullets = [bullets]
# #         elif not isinstance(bullets, list):
# #             bullets = []
# #         notes = parsed.get("notes", "")
# #         return {"title": str(title), "bullets": [str(b) for b in bullets], "notes": str(notes)}
# # 
# #     # Fallback template slide
# #     return {
# #         "title": topic.title(),
# #         "bullets": [
# #             f"Overview of {topic}",
# #             "Key definitions and background concepts",
# #             "Primary challenges and modern implications",
# #         ],
# #         "notes": f"This slide presents key details about {topic}.",
# #     }
# # 
# # 
# # async def ppt_add_slide(args: dict, session_id: str) -> dict:
# #     from backend.api.ppt import save_and_render_pptx
# #     from backend.core.slide_store import resolve_sid
# #     from backend.queues.bus import bus
# # 
# #     effective_sid = resolve_sid(session_id)
# #     slides = list(_slide_store.get(effective_sid, []))
# # 
# #     topic = args.get("topic", "").strip(" .?!")
# #     if not topic:
# #         topic = "New Slide"
# # 
# #     logger.info(f"Adding a slide on topic '{topic}' via voice...")
# # 
# #     # Generate content
# #     if topic != "New Slide":
# #         new_slide = await _generate_single_slide_content(topic)
# #     else:
# #         new_slide = {
# #             "title": "New Slide",
# #             "bullets": ["Write content or use voice to edit this slide"],
# #             "notes": "",
# #         }
# # 
# #     # Recompile plain slide data preserving styling and images
# #     slides_to_save = _prepare_slides_for_save(slides)
# #     slides_to_save.append(new_slide)
# # 
# #     # Save and render
# #     rendered = save_and_render_pptx(slides_to_save, effective_sid, topic=topic)
# #     new_idx = len(rendered) - 1
# #     _current_slide[effective_sid] = new_idx
# # 
# #     await bus.emit_event(
# #         "ppt_command",
# #         {
# #             "action": "reload",
# #             "slides": rendered,
# #             "filename": f"Presentation ({len(rendered)} slides)",
# #             "index": new_idx,
# #             "preserveCurrent": False,
# #         },
# #         session_id,
# #     )
# # 
# #     return {
# #         "status": "ok",
# #         "spoken_reply": f"I've added a slide about {new_slide['title']} and navigated to it.",
# #     }
# # 
# # 
# # async def ppt_edit_slide(args: dict, session_id: str) -> dict:
# #     from backend.api.ppt import save_and_render_pptx
# #     from backend.core.slide_store import resolve_sid
# #     from backend.queues.bus import bus
# # 
# #     effective_sid = resolve_sid(session_id)
# #     slides = list(_slide_store.get(effective_sid, []))
# #     if not slides:
# #         return {"status": "error", "spoken_reply": "There is no active presentation to edit."}
# # 
# #     current_idx = _current_slide.get(effective_sid, 0)
# #     if current_idx < 0 or current_idx >= len(slides):
# #         current_idx = 0
# # 
# #     title = args.get("title")
# #     bullets = args.get("bullets")
# #     add_bullet = args.get("add_bullet")
# #     notes = args.get("notes")
# # 
# #     # Font style properties
# #     title_bold = args.get("title_bold")
# #     title_italic = args.get("title_italic")
# #     bullet_bold = args.get("bullet_bold")
# #     bullet_italic = args.get("bullet_italic")
# #     font_name = args.get("font_name")
# # 
# #     spoken_reply = "Updated the slide successfully."
# # 
# #     if title is not None:
# #         slides[current_idx]["title"] = title
# #         spoken_reply = f"I have updated the slide title to '{title}'."
# #     if bullets is not None:
# #         slides[current_idx]["bullets"] = bullets
# #         spoken_reply = "I have updated the bullet points on this slide."
# #     elif add_bullet is not None:
# #         if "bullets" not in slides[current_idx] or not isinstance(slides[current_idx]["bullets"], list):
# #             slides[current_idx]["bullets"] = []
# #         slides[current_idx]["bullets"].append(add_bullet)
# #         spoken_reply = f"I've added a bullet point about '{add_bullet}' to this slide."
# #     if notes is not None:
# #         slides[current_idx]["notes"] = notes
# #         spoken_reply = "I have updated the speaker notes for this slide."
# # 
# #     # Apply styling properties
# #     style_updated = []
# #     if title_bold is not None:
# #         slides[current_idx]["title_bold"] = title_bold
# #         style_updated.append("bold" if title_bold else "normal")
# #     if title_italic is not None:
# #         slides[current_idx]["title_italic"] = title_italic
# #         style_updated.append("italic" if title_italic else "normal")
# #     if bullet_bold is not None:
# #         slides[current_idx]["bullet_bold"] = bullet_bold
# #         style_updated.append("bold bullets" if bullet_bold else "normal bullets")
# #     if bullet_italic is not None:
# #         slides[current_idx]["bullet_italic"] = bullet_italic
# #         style_updated.append("italic bullets" if bullet_italic else "normal bullets")
# #     if font_name is not None:
# #         slides[current_idx]["font_name"] = font_name
# #         style_updated.append(f"font to {font_name}")
# # 
# #     if style_updated:
# #         spoken_reply = f"I have updated the slide formatting: {', '.join(style_updated)}."
# # 
# #     # Recompile plain slide data preserving images and styling details
# #     slides_to_save = _prepare_slides_for_save(slides)
# # 
# #     rendered = save_and_render_pptx(slides_to_save, effective_sid, topic="Presentation")
# #     _current_slide[effective_sid] = current_idx
# # 
# #     await bus.emit_event(
# #         "ppt_command",
# #         {
# #             "action": "reload",
# #             "slides": rendered,
# #             "filename": f"Presentation ({len(rendered)} slides)",
# #             "index": current_idx,
# #             "preserveCurrent": True,
# #         },
# #         session_id,
# #     )
# # 
# #     return {"status": "ok", "spoken_reply": spoken_reply}
# # 
# # 
# 
# async def ppt_clear_presentation(args: dict, session_id: str) -> dict:
#     from backend.api.ppt import save_and_render_pptx
#     from backend.core.slide_store import resolve_sid
#     from backend.queues.bus import bus
# 
#     effective_sid = resolve_sid(session_id)
#     default_slides = [
#         {
#             "title": "New Presentation",
#             "bullets": ["Voice-Driven Interactive Presentation", "Start speaking or clicking to add slides"],
#             "notes": "Welcome to your new interactive presentation deck.",
#         }
#     ]
# 
#     rendered = save_and_render_pptx(default_slides, effective_sid, topic="New Presentation")
#     _current_slide[effective_sid] = 0
# 
#     await bus.emit_event(
#         "ppt_command",
#         {
#             "action": "reload",
#             "slides": rendered,
#             "filename": "New Presentation",
#             "index": 0,
#             "preserveCurrent": False,
#         },
#         session_id,
#     )
# 
#     return {
#         "status": "ok",
#         "spoken_reply": "I have started a new presentation deck with a blank title slide. What topic would you like to add a slide about?",
#     }
# 
# 
# async def ppt_improvise_slide(args: dict, session_id: str) -> dict:
#     from backend.api.ppt import improvise_slide_content, save_and_render_pptx
#     from backend.core.slide_store import resolve_sid
#     from backend.queues.bus import bus
# 
#     effective_sid = resolve_sid(session_id)
#     slides = list(_slide_store.get(effective_sid, []))
#     if not slides:
#         return {"status": "error", "spoken_reply": "There is no active presentation to improvise."}
# 
#     current_idx = _current_slide.get(effective_sid, 0)
#     if current_idx < 0 or current_idx >= len(slides):
#         current_idx = 0
# 
#     prompt = args.get("prompt", "").strip(" .?!")
#     if not prompt:
#         return {"status": "error", "spoken_reply": "What would you like me to improve on this slide?"}
# 
#     logger.info(f"Improvising slide {current_idx + 1} with prompt: '{prompt}' via voice...")
# 
#     try:
#         current_slide = slides[current_idx]
#         improved = await improvise_slide_content(current_slide, prompt)
# 
#         # Update text fields but preserve any images!
#         slides[current_idx]["title"] = improved["title"]
#         slides[current_idx]["bullets"] = improved["bullets"]
#         slides[current_idx]["notes"] = improved["notes"]
# 
#         # Recompile plain slide data preserving images and style details
#         slides_to_save = _prepare_slides_for_save(slides)
# 
#         rendered = save_and_render_pptx(slides_to_save, effective_sid, topic="Presentation")
#         _current_slide[effective_sid] = current_idx
# 
#         await bus.emit_event(
#             "ppt_command",
#             {
#                 "action": "reload",
#                 "slides": rendered,
#                 "filename": f"Presentation ({len(rendered)} slides)",
#                 "index": current_idx,
#                 "preserveCurrent": True,
#             },
#             session_id,
#         )
# 
#         return {
#             "status": "ok",
#             "spoken_reply": f"I've improvised this slide according to your request: {prompt}.",
#         }
#     except Exception as e:
#         logger.error(f"Error improvising slide via voice: {e}")
#         return {
#             "status": "error",
#             "spoken_reply": "I encountered an error trying to improvise the slide. Please try again.",
#         }
# 
# 
# async def ppt_save_slide(args: dict, session_id: str) -> dict:
#     """Triggers the frontend to save any unsaved slide edits (title, bullets, notes) via event bus."""
#     from backend.queues.bus import bus
#     
#     logger.info("Triggering save slide changes command via voice event...")
#     await bus.emit_event(
#         "ppt_command",
#         {
#             "action": "save"
#         },
#         session_id
#     )
#     return {
#         "status": "ok",
#         "spoken_reply": "I have saved your slide changes."
#     }
# 
# 
# async def route_page(args: dict, session_id: str) -> dict:
#     """Navigates the frontend page dynamically to the specified target route via voice."""
#     from backend.queues.bus import bus
#     
#     page = args.get("page", "dashboard")
#     logger.info(f"Triggering page route transition via voice: target={page}")
#     await bus.emit_event(
#         "route_page",
#         {
#             "page": page
#         },
#         session_id
#     )
#     name_map = {
#         "care": "Trip Planner",
#         "ppt": "PPT Copilot",
#         "email": "Email Center",
#         "meetings": "Meeting TALKINIA",
#         "chat": "Wanna Chat",
#         "guideline": "System Guidelines",
#         "profile": "Profile Page",
#         "settings": "Settings Page",
#         "dashboard": "Main Dashboard"
#     }
#     pretty_name = name_map.get(page, "Main Dashboard")
#     return {
#         "status": "ok",
#         "spoken_reply": f"Navigating to {pretty_name}."
#     }
# 
# # ============================================================================
# # ACTIVE — GD-template kind-aware slide editing, notes generation, add-slide.
# # Migrated from syugesh/pilot-voice-agent-backend (feature/ppt-copilot @
# # 2d3df73). Supersedes ppt_navigate / ppt_jump_to_title / ppt_delete_slide /
# # ppt_summarize / ppt_add_slide / ppt_edit_slide above. PILOT-native
# # ppt_create_slides / ppt_clear_presentation / ppt_improvise_slide / ppt_qa /
# # ppt_save_slide / route_page (kept above) have no upstream equivalent.
# # ============================================================================
# """PPT tools — navigate + jump to slide by number or title + summarize."""
# import asyncio, logging
# logger = logging.getLogger("pilot.tools.ppt")
# 
# 
# async def ppt_navigate(args: dict, session_id: str) -> dict:
#     direction = args.get("direction", "next")
#     from backend.queues.bus import bus
#     from backend.api.ppt import _slide_store, get_latest_upload_sid, _current_slide
# 
#     slides = _slide_store.get(session_id) or _slide_store.get(get_latest_upload_sid(), [])
#     effective_sid = session_id if session_id in _slide_store else get_latest_upload_sid()
#     total = len(slides)
#     current = _current_slide.get(effective_sid, 0)
# 
#     if direction == "next":
#         if total > 0 and current >= total - 1:
#             return {"spoken_reply": f"You've reached the last slide — slide {total} of {total}. That's the end of the presentation."}
#         _current_slide[effective_sid] = min(current + 1, max(total - 1, 0))
#     elif direction == "prev":
#         if current <= 0:
#             return {"spoken_reply": "You're already on the first slide."}
#         _current_slide[effective_sid] = max(current - 1, 0)
#     elif direction == "first":
#         _current_slide[effective_sid] = 0
#     elif direction == "last":
#         _current_slide[effective_sid] = max(total - 1, 0)
# 
#     new_idx = _current_slide.get(effective_sid, 0)
#     await bus.emit_event("ppt_command", {"action": direction}, session_id)
#     return {"status": "ok", "direction": direction, "index": new_idx, "spoken_reply": ""}
# 
# 
# async def ppt_jump_to_title(args: dict, session_id: str) -> dict:
#     query        = args.get("query", "")
#     slide_number = args.get("slide_number")   # already 0-indexed if from keyword fallback
#     from backend.queues.bus import bus
# 
#     from backend.api.ppt import _slide_store, get_latest_upload_sid, _current_slide
#     import re
# 
#     effective_sid = session_id if session_id in _slide_store else get_latest_upload_sid()
#     slides = _slide_store.get(effective_sid, [])
# 
#     # If explicit slide number given, use directly
#     if slide_number is not None:
#         idx = int(slide_number)
#         _current_slide[effective_sid] = idx
#         await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
#         return {"status": "ok", "index": idx, "title": f"Slide {idx+1}", "spoken_reply": ""}
# 
#     q = query.lower()
# 
#     # Numeric match in query — only allow within actual deck bounds
#     m = re.search(r'\b(\d+)\b', q)
#     if m:
#         idx = int(m.group(1)) - 1
#         if 0 <= idx < len(slides):
#             _current_slide[effective_sid] = idx
#             await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
#             return {"status": "ok", "index": idx, "spoken_reply": ""}
# 
#     # Title fuzzy match
#     best_idx, best_score = 0, 0
#     for s in slides:
#         score = sum(1 for w in q.split() if w in s.get("title","").lower())
#         if score > best_score:
#             best_score = score; best_idx = s["index"]
# 
#     _current_slide[effective_sid] = best_idx
#     await bus.emit_event("ppt_command", {"action": "goto", "index": best_idx}, session_id)
#     return {"status": "ok", "index": best_idx, "spoken_reply": ""}
# 
# 
# async def ppt_delete_slide(args: dict, session_id: str) -> dict:
#     from backend.queues.bus import bus
#     slide_number = args.get("slide_number")
#     if slide_number is not None:
#         # Navigate to the target slide first so the confirm modal shows the right one
#         await bus.emit_event("ppt_command", {"action": "goto", "index": int(slide_number)}, session_id)
#     await bus.emit_event("ppt_command", {"action": "delete"}, session_id)
#     label = f"slide {int(slide_number) + 1}" if slide_number is not None else "this slide"
#     return {"status": "ok", "spoken_reply": f"Please confirm the deletion of {label} in the popup."}
# 
# 
# async def ppt_summarize(args: dict, session_id: str) -> dict:
#     from backend.api.ppt import _slide_store, get_latest_upload_sid
#     slides = _slide_store.get(session_id) or _slide_store.get(get_latest_upload_sid(), [])
#     if not slides:
#         return {"spoken_reply": "No presentation is loaded yet. Please upload a PowerPoint file first."}
# 
#     lines = []
#     for s in slides:
#         title = s.get("title", f"Slide {s['index']+1}")
#         notes = s.get("notes", "")
#         lines.append(f"Slide {s['index']+1}: {title}" + (f" — {notes}" if notes else ""))
# 
#     content = "\n".join(lines)
#     summary = await _summarize_ollama(content)
#     return {"spoken_reply": summary}
# 
# 
# async def ppt_last_action(args: dict, session_id: str) -> dict:
#     from backend.core.session_state import get_state
#     state = get_state(session_id)
#     last = state.last_ppt_action or {}
#     if not last:
#         return {"spoken_reply": "I haven't changed the deck yet in this session."}
#     slide_no = last.get("slide_number")
#     instruction = last.get("instruction", "your last request")
#     changes = last.get("changes") or []
#     if changes:
#         changed_text = ", ".join(changes)
#         return {"spoken_reply": f"Yes. I updated slide {slide_no}: {changed_text}, based on your request to {instruction}."}
#     return {"spoken_reply": f"I tried to update slide {slide_no} based on your request to {instruction}, but I don't see a content change."}
# 
# 
# async def _summarize_ollama(content: str) -> str:
#     def _call() -> str:
#         import ollama
#         from backend.core.config import settings
#         resp = ollama.chat(
#             model=settings.OLLAMA_MODEL,
#             messages=[
#                 {"role": "system", "content":
#                     "You are a helpful voice assistant. Summarize the presentation in 4-6 natural spoken sentences. "
#                     "Mention the main topics and key points. No markdown, no bullet points — plain conversational speech only."},
#                 {"role": "user", "content": f"Summarize this presentation:\n\n{content[:3000]}"},
#             ],
#             think=False,
#             options={"num_predict": 220},
#             stream=False,
#         )
#         if isinstance(resp, dict):
#             return resp["message"]["content"].strip()
#         return resp.message.content.strip()
# 
#     try:
#         return await asyncio.to_thread(_call)
#     except Exception as e:
#         logger.error(f"ppt_summarize ollama error: {e}")
#         return "I wasn't able to summarize the presentation right now. Please try again."
# 
# 
# async def _apply_edit_and_background_refresh(effective_sid: str, idx: int, session_id: str, writer) -> bool:
#     """
#     Writes an edit immediately (fast — well under a second) so the caller
#     can speak a confirmation right away, then regenerates thumbnails
#     (LibreOffice re-rendering the whole deck — several seconds) as a
#     background task, emitting the same "refresh" ppt_command once that
#     finishes instead of making the voice reply wait for it too. The
#     thumbnail visually catches up a couple seconds after the spoken
#     confirmation — the frontend already refetches on that event regardless
#     of when it fires.
#     Returns True on success, False if the session/slide index was invalid.
#     """
#     from backend.api.ppt import apply_slide_edit_fast_async, refresh_slide_thumbnails_async
#     from backend.queues.bus import bus
# 
#     ok = await apply_slide_edit_fast_async(effective_sid, idx, writer)
#     if not ok:
#         return False
# 
#     async def _background_refresh():
#         await refresh_slide_thumbnails_async(effective_sid)
#         await bus.emit_event("ppt_command", {"action": "refresh", "session_id": effective_sid}, session_id)
# 
#     asyncio.create_task(_background_refresh())
#     return True
# 
# 
# async def ppt_edit_slide(args: dict, session_id: str) -> dict:
#     instruction = args.get("instruction", "")
#     if not instruction:
#         return {"spoken_reply": "What would you like me to change?"}
# 
#     from backend.api.ppt import _slide_store, get_latest_upload_sid, _current_slide
#     from backend.queues.bus import bus
#     import re
# 
#     effective_sid = session_id if session_id in _slide_store else get_latest_upload_sid()
#     slides = _slide_store.get(effective_sid, [])
#     if not slides:
#         return {"spoken_reply": "No presentation is loaded."}
#         
#     slide_number = args.get("slide_number")
#     idx = int(slide_number) if slide_number is not None else _current_slide.get(effective_sid, 0)
#     if idx < 0 or idx >= len(slides):
#         return {"spoken_reply": "I'm not sure which slide to edit."}
#     _current_slide[effective_sid] = idx
#     await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
# 
#     slide = slides[idx]
# 
#     # Slides generated from the GD template have a known "kind" (team,
#     # table, comparison, ...) whose content doesn't fit the generic
#     # title/numbered-bullets model at all — route those through the
#     # kind-aware editor instead. Uploaded/legacy decks have no kind and
#     # fall straight through to the unchanged generic path below.
#     kind = slide.get("kind")
#     if kind:
#         return await _edit_kind_aware_slide(effective_sid, idx, kind, instruction, session_id)
# 
#     # Extract bullets
#     old_bullets = []
#     for sh in slide.get("shapes", []):
#         for line in sh.get("text", "").split("\n"):
#             m = re.match(r'^\d+\.\s{1,3}(.+)', line)
#             if m:
#                 old_bullets.append(m.group(1).strip())
#                 
#     old_title = slide.get("title", "")
#     old_notes = slide.get("notes", "")
# 
#     # Every other text-bearing shape that ISN'T the title or a numbered bullet
#     # line (e.g. a subtitle, caption, or footer). Keyed by stable shape_id so
#     # the editor can target one exact shape without touching anything else on
#     # the slide — this is what lets the user edit "anything on the slide",
#     # not just the title/bullets/notes triplet.
#     other_shapes: dict[str, str] = {}
#     for sh in slide.get("shapes", []):
#         sid = sh.get("shape_id")
#         text = sh.get("text", "").strip()
#         if sid is None or not text:
#             continue
#         if text == old_title:
#             continue
#         if all(re.match(r'^\d+\.\s{1,3}', line) for line in text.split("\n") if line.strip()):
#             continue  # this is the bullet body shape, already covered above
#         other_shapes[str(sid)] = text
# 
#     # Common voice correction: "change the title of slide 3 from X to Y".
#     # Handle it deterministically so ASR/correction commands don't depend on
#     # the editor LLM preserving the exact replacement text.
#     title_to = re.search(r'\btitle\b.*?\bfrom\b.+?\bto\b\s+(.+)$', instruction, re.IGNORECASE)
#     title_direct = re.search(r'\b(?:change|update|set|rename)\b.*?\btitle\b.*?\bto\b\s+(.+)$', instruction, re.IGNORECASE)
#     replacement_title = (title_to or title_direct)
#     if replacement_title:
#         new_title = replacement_title.group(1).strip().strip('".,!?')
#         if new_title:
#             try:
#                 from backend.api.ppt import _patch_slide
#                 # other_edits intentionally omitted (defaults to None): a
#                 # title-only instruction must never touch bullets, notes, or
#                 # any other shape on the slide.
#                 def _writer(slide, _t=new_title, _b=old_bullets, _n=old_notes):
#                     _patch_slide(slide, idx, _t, _b, _n, None)
#                 await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
#                 from backend.core.session_state import get_state
#                 state = get_state(session_id)
#                 state.last_ppt_action = {
#                     "tool": "ppt_edit_slide",
#                     "slide_number": idx + 1,
#                     "instruction": instruction,
#                     "changes": [f"title from '{old_title}' to '{new_title}'"],
#                 }
#                 return {"spoken_reply": f"I changed the title on slide {idx + 1} to {new_title}."}
#             except Exception as e:
#                 logger.error(f"ppt_edit_slide title shortcut error: {e}")
#                 return {"spoken_reply": "I had trouble editing the slide right now."}
# 
#     # LLM call
#     def _call() -> str:
#         import ollama
#         from backend.core.config import settings
#         import json
#         other_shapes_desc = (
#             "\n".join(f'- shape_id {sid}: "{text}"' for sid, text in other_shapes.items())
#             if other_shapes else "(none)"
#         )
#         prompt = f"""
# You are an expert presentation editor. The user wants to edit a slide.
# Current Slide:
# Title: {old_title}
# Bullets: {json.dumps(old_bullets)}
# Notes: {old_notes}
# Other text on this slide (subtitle/caption/footer — NOT the title or bullets):
# {other_shapes_desc}
# 
# Instruction: {instruction}
# 
# Rules:
# - Only change what the instruction actually asks for. Fields you are not asked
#   to change must be echoed back EXACTLY as given above — do not paraphrase,
#   shorten, or drop unrelated text.
# - If the instruction targets one of the "other text" shapes above (a subtitle,
#   caption, footer, etc.), put its new text in "other_edits" keyed by its
#   shape_id. Do not invent shape_ids that weren't listed. Leave "other_edits"
#   as {{}} if nothing else needs to change.
# 
# Output ONLY valid JSON matching this schema:
# {{
#   "title": "New Title",
#   "bullets": ["New Bullet 1", "New Bullet 2"],
#   "notes": "New Speaker Notes",
#   "other_edits": {{"<shape_id>": "new text for that shape"}}
# }}
# """
#         resp = ollama.chat(
#             model=settings.OLLAMA_MODEL,
#             messages=[{"role": "user", "content": prompt}],
#             # think=False is required alongside format="json" — without it, a
#             # reasoning-capable model (Qwen3, etc.) tries to emit chain-of-thought
#             # before its JSON, which conflicts with grammar-constrained JSON
#             # decoding and can come back as empty content (see _summarize_ollama
#             # above, and _generate_template_content_sync in
#             # services/ppt_template_builder.py, for the same fix).
#             options={"num_predict": 400},
#             format="json",
#             think=False,
#             stream=False,
#         )
#         if isinstance(resp, dict):
#             return resp["message"]["content"].strip()
#         return resp.message.content.strip()
# 
#     try:
#         raw = await asyncio.to_thread(_call)
#         import json
#         if not raw:
#             # Same empty-content failure mode this fix targets — retry once
#             # before giving up, since Ollama occasionally still returns a
#             # blank message.strip() on a cold model load.
#             raw = await asyncio.to_thread(_call)
#         if not raw:
#             logger.error("ppt_edit_slide error: Ollama returned an empty response after retry")
#             return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}
#         try:
#             data = json.loads(raw)
#         except json.JSONDecodeError:
#             logger.error(f"ppt_edit_slide error: invalid JSON from Ollama: {raw[:300]!r}")
#             return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}
#         new_title = data.get("title", old_title)
#         new_bullets = data.get("bullets", old_bullets)
#         new_notes = data.get("notes", old_notes)
# 
#         # Only accept other_edits for shape_ids we actually listed to the
#         # model — guards against a hallucinated id landing on the wrong shape.
#         raw_other_edits = data.get("other_edits") or {}
#         other_edits = {
#             sid: text for sid, text in raw_other_edits.items()
#             if sid in other_shapes and isinstance(text, str) and text.strip()
#         }
# 
#         changes = []
#         if new_title != old_title:
#             changes.append(f"title from '{old_title}' to '{new_title}'")
#         if new_bullets != old_bullets:
#             changes.append("bullet content")
#         if new_notes != old_notes:
#             changes.append("speaker notes")
#         for sid, text in other_edits.items():
#             if text != other_shapes.get(sid):
#                 changes.append("other slide text")
#                 break
#         
#         from backend.api.ppt import _patch_slide
#         def _writer(slide, _t=new_title, _b=new_bullets, _n=new_notes, _oe=other_edits or None):
#             _patch_slide(slide, idx, _t, _b, _n, _oe)
#         await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
#         from backend.core.session_state import get_state
#         state = get_state(session_id)
#         state.last_ppt_action = {
#             "tool": "ppt_edit_slide",
#             "slide_number": idx + 1,
#             "instruction": instruction,
#             "changes": changes,
#         }
#         return {"spoken_reply": f"I've updated slide {idx + 1}."}
#     except Exception as e:
#         logger.error(f"ppt_edit_slide error: {e}")
#         return {"spoken_reply": "I had trouble editing the slide right now."}
# 
# 
# async def _edit_kind_aware_slide(effective_sid: str, idx: int, kind: str, instruction: str, session_id: str) -> dict:
#     """
#     Voice-edit path for slides cloned from the GD template (team, table,
#     comparison, ...) — the generic title/numbered-bullets model above
#     doesn't fit their structure at all. Same "send current state as JSON,
#     ask the model for edited JSON back" pattern as the generic path, just
#     with each kind's own field schema instead of title/bullets/notes.
#     """
#     from backend.services.ppt_template_builder import extract_slide_data, populate_slide_data
#     import json
# 
#     def _read_current() -> dict:
#         from pptx import Presentation
#         prs = Presentation(f"data/ppt/{effective_sid}.pptx")
#         return extract_slide_data(prs.slides[idx], kind)
# 
#     try:
#         current_data = await asyncio.to_thread(_read_current)
#     except Exception as e:
#         logger.error(f"kind-aware edit: failed to read current slide data: {e}")
#         return {"spoken_reply": "I had trouble reading that slide right now."}
# 
#     def _call() -> str:
#         import ollama
#         from backend.core.config import settings
#         prompt = f"""
# You are an expert presentation editor. The user wants to edit a "{kind}" slide.
# 
# Current content (JSON): {json.dumps(current_data)}
# 
# Instruction: {instruction}
# 
# Rules:
# - Return the SAME JSON shape as "Current content" above, with the requested
#   change applied.
# - Only change what the instruction actually asks for — every other field
#   must be echoed back EXACTLY as given, do not paraphrase or drop anything.
# - Output ONLY valid JSON — no markdown, no explanation, nothing else.
# """
#         resp = ollama.chat(
#             model=settings.OLLAMA_MODEL,
#             messages=[{"role": "user", "content": prompt}],
#             options={"num_predict": 500},
#             format="json",
#             think=False,  # see ppt_edit_slide above for why this must pair with format="json"
#             stream=False,
#         )
#         if isinstance(resp, dict):
#             return resp["message"]["content"].strip()
#         return resp.message.content.strip()
# 
#     try:
#         raw = await asyncio.to_thread(_call)
#         if not raw:
#             raw = await asyncio.to_thread(_call)  # cold-model-load retry, same as the generic path
#         if not raw:
#             logger.error("kind-aware edit error: Ollama returned an empty response after retry")
#             return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}
#         try:
#             new_data = json.loads(raw)
#         except json.JSONDecodeError:
#             logger.error(f"kind-aware edit error: invalid JSON from Ollama: {raw[:300]!r}")
#             return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}
# 
#         def _writer(slide):
#             populate_slide_data(slide, kind, new_data)
# 
#         await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
#         from backend.core.session_state import get_state
#         state = get_state(session_id)
#         state.last_ppt_action = {
#             "tool": "ppt_edit_slide",
#             "slide_number": idx + 1,
#             "instruction": instruction,
#             "changes": ["slide content"],
#         }
#         return {"spoken_reply": f"I've updated slide {idx + 1}."}
#     except Exception as e:
#         logger.error(f"kind-aware ppt_edit_slide error: {e}")
#         return {"spoken_reply": "I had trouble editing the slide right now."}
# 
# 
# def _slide_bullets_for_notes(slide: dict) -> list:
#     import re
#     bullets = []
#     for sh in slide.get("shapes", []):
#         for line in sh.get("text", "").split("\n"):
#             m = re.match(r'^\d+\.\s{1,3}(.+)', line)
#             if m:
#                 bullets.append(m.group(1).strip())
#     return bullets
# 
# 
# def _generate_notes_text_sync(title: str, bullets: list) -> str:
#     import ollama, json
#     from backend.core.config import settings
#     prompt = f"""
# Write 3-4 natural conversational sentences of speaker notes for this slide.
# Title: {title}
# Bullets: {json.dumps(bullets)}
# 
# Output the speaker notes plainly without any markdown, prefix, or JSON formatting.
# """
#     resp = ollama.chat(
#         model=settings.OLLAMA_MODEL,
#         messages=[{"role": "user", "content": prompt}],
#         options={"num_predict": 250},
#         stream=False,
#     )
#     if isinstance(resp, dict):
#         return resp["message"]["content"].strip()
#     return resp.message.content.strip()
# 
# 
# async def ppt_generate_notes(args: dict, session_id: str) -> dict:
#     from backend.api.ppt import (
#         _slide_store, get_latest_upload_sid, _current_slide,
#         apply_notes_batch_async,
#     )
#     from backend.queues.bus import bus
# 
#     effective_sid = session_id if session_id in _slide_store else get_latest_upload_sid()
#     slides = _slide_store.get(effective_sid, [])
#     if not slides:
#         return {"spoken_reply": "No presentation is loaded."}
# 
#     if args.get("all"):
#         await bus.emit_event("ppt_command", {"action": "goto", "index": 0}, session_id)
#         notes_by_index: dict[int, str] = {}
#         failures = 0
#         for i, slide in enumerate(slides):
#             try:
#                 notes_by_index[i] = await asyncio.to_thread(
#                     _generate_notes_text_sync, slide.get("title", ""), _slide_bullets_for_notes(slide),
#                 )
#             except Exception as e:
#                 failures += 1
#                 logger.error(f"ppt_generate_notes (all) slide {i} error: {e}")
#         if not notes_by_index:
#             return {"spoken_reply": "I had trouble generating notes right now."}
#         await apply_notes_batch_async(effective_sid, notes_by_index)
#         await bus.emit_event("ppt_command", {"action": "refresh", "session_id": effective_sid}, session_id)
#         from backend.core.session_state import get_state
#         state = get_state(session_id)
#         state.last_ppt_action = {
#             "tool": "ppt_generate_notes",
#             "slide_number": None,
#             "instruction": "generate speaker notes for all slides",
#             "changes": [f"speaker notes for {len(notes_by_index)} of {len(slides)} slides"],
#         }
#         if failures:
#             return {"spoken_reply": f"I've generated speaker notes for {len(notes_by_index)} of {len(slides)} slides — {failures} failed."}
#         return {"spoken_reply": f"I've generated speaker notes for all {len(slides)} slides."}
# 
#     slide_number = args.get("slide_number")
#     idx = int(slide_number) if slide_number is not None else _current_slide.get(effective_sid, 0)
#     if idx < 0 or idx >= len(slides):
#         return {"spoken_reply": "I'm not sure which slide to write notes for."}
#     _current_slide[effective_sid] = idx
#     await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
#     slide = slides[idx]
#     old_title = slide.get("title", "")
#     old_bullets = _slide_bullets_for_notes(slide)
# 
#     try:
#         notes = await asyncio.to_thread(_generate_notes_text_sync, old_title, old_bullets)
#         # Apply notes, keep title and bullets the same
#         from backend.api.ppt import _patch_slide
#         def _writer(slide, _t=old_title, _b=old_bullets, _n=notes):
#             _patch_slide(slide, idx, _t, _b, _n, None)
#         await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
#         from backend.core.session_state import get_state
#         state = get_state(session_id)
#         state.last_ppt_action = {
#             "tool": "ppt_generate_notes",
#             "slide_number": idx + 1,
#             "instruction": "generate speaker notes",
#             "changes": ["speaker notes"],
#         }
#         return {"spoken_reply": f"I've generated new speaker notes for slide {idx + 1}."}
#     except Exception as e:
#         logger.error(f"ppt_generate_notes error: {e}")
#         return {"spoken_reply": "I had trouble generating notes right now."}
# 
# 
# async def ppt_add_slide(args: dict, session_id: str) -> dict:
#     """Insert a new slide into the loaded presentation, content generated
#     from a spoken description ("add a slide about our Q4 roadmap"). Falls
#     back to a plain title-only text slide if generation fails or no
#     description was given, rather than failing the whole command."""
#     from backend.api.ppt import _slide_store, get_latest_upload_sid
#     from backend.services.ppt_template_builder import generate_single_slide_content
# 
#     effective_sid = session_id if session_id in _slide_store else get_latest_upload_sid()
#     slides = _slide_store.get(effective_sid, [])
#     if not slides:
#         return {"spoken_reply": "No presentation is loaded yet. Please upload or create one first."}
# 
#     instruction = args.get("instruction", "").strip()
#     slide_data = None
#     if instruction:
#         try:
#             slide_data = await generate_single_slide_content(instruction)
#         except Exception as e:
#             logger.error(f"ppt_add_slide generation error: {e}")
# 
#     if not slide_data:
#         title = instruction[:60] if instruction else "New Slide"
#         slide_data = {"kind": "text", "title": title, "paragraphs": []}
# 
#     kind = slide_data.pop("kind")
# 
#     try:
#         ok = await _apply_add_and_background_refresh(effective_sid, kind, slide_data, session_id)
#         if not ok:
#             return {"spoken_reply": "I couldn't add a slide right now."}
#         from backend.core.session_state import get_state
#         state = get_state(session_id)
#         state.last_ppt_action = {
#             "tool": "ppt_add_slide",
#             "slide_number": len(slides) + 1,
#             "instruction": instruction or "add a new slide",
#             "changes": [f"new {kind} slide"],
#         }
#         return {"spoken_reply": f"I've added a new slide — {slide_data.get('title', 'untitled')}."}
#     except Exception as e:
#         logger.error(f"ppt_add_slide error: {e}")
#         return {"spoken_reply": "I had trouble adding that slide."}
# 
# 
# async def _apply_add_and_background_refresh(effective_sid: str, kind: str, data: dict, session_id: str) -> bool:
#     """Same fast-reply-then-background-refresh pattern as
#     _apply_edit_and_background_refresh, for the add-slide path: writes the
#     new slide immediately, then regenerates thumbnails (LibreOffice, whole
#     deck) in the background and re-emits ppt_command so the viewer catches
#     up a moment after the spoken confirmation instead of before it."""
#     from backend.api.ppt import add_slide_fast_async, refresh_slide_thumbnails_async
#     from backend.queues.bus import bus
# 
#     new_index = await add_slide_fast_async(effective_sid, kind, data)
#     if new_index is None:
#         return False
# 
#     async def _background_refresh():
#         await refresh_slide_thumbnails_async(effective_sid)
#         await bus.emit_event("ppt_command", {"action": "goto", "index": new_index}, session_id)
#         await bus.emit_event("ppt_command", {"action": "refresh", "session_id": effective_sid}, session_id)
# 
#     asyncio.create_task(_background_refresh())
#     return True
# ============================================================================
# DISABLED: ppt copilot implementation transferred from
# syugesh/pilot-voice-agent-backend (feature/ppt-copilot @ commit 869dd08d9).
# Superseded by the newer pilot-voice-agent-backend-ppt_update snapshot,
# pasted in active below. Kept here commented out for reference only.
# ============================================================================
# """PPT tools — navigate + jump to slide by number or title + summarize."""
# # (updated from feature/ppt-copilot @ c3a22cfc3 — adds ppt_reorder_slide,
# # smarter insert-position parsing, auto-positioning, GD-kind hardening)
# """PPT tools — navigate + jump to slide by number or title + summarize."""
# import asyncio, logging, re
# logger = logging.getLogger("pilot.tools.ppt")
#
#
# # ── Add-slide helpers — parse "where" out of the instruction, and detect
# # when "what" was never actually said (so we ask instead of hallucinating
# # generic filler content that only coincidentally resembles a real slide) ──
#
# _ORDINALS = {
#     "first": 1, "second": 2, "third": 3, "fourth": 4, "fifth": 5,
#     "sixth": 6, "seventh": 7, "eighth": 8, "ninth": 9, "tenth": 10,
# }
#
# def _parse_insert_position(text: str) -> tuple[int | None, bool]:
#     """Returns (insert_after, was_specified). insert_after is 0-indexed —
#     the new slide lands right after prs.slides[insert_after]. -1 means "at the
#     very beginning" (before slide 1), None means "append at the end" (also the
#     default when nothing was said)."""
#     t = text.lower()
#
#     # "after slide N" → land right after slide N (0-indexed N-1).
#     m = re.search(r'\bafter\s+slide\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
#     if m:
#         return int(m.group(1)) - 1, True
#
#     # "before slide N" → before slide N == after slide N-1 (0-indexed N-2).
#     m = re.search(r'\bbefore\s+slide\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
#     if m:
#         return int(m.group(1)) - 2, True
#
#     # "on/at slide N", "at position N", "as slide N", "make it slide N",
#     # "in position N", or a bare "slide N" — all mean "the new slide should
#     # BECOME slide N", i.e. occupy position N. That's after slide N-1 (0-indexed
#     # N-2), so slide N-1 stays before it and the old slide N shifts down.
#     m = re.search(
#         r'\b(?:on|at|as|in(?:to)?|position|make\s+it)\s+(?:the\s+)?'
#         r'(?:slide|position)?\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
#     if m:
#         n = int(m.group(1))
#         return (-1 if n <= 1 else n - 2), True
#
#     # Ordinal word forms: "as the third slide", "make it the second slide".
#     m = re.search(r'\b(?:as|make\s+it|at|position)\s+(?:the\s+)?(' + "|".join(_ORDINALS) + r')\s+slide\b', t)
#     if m:
#         n = _ORDINALS[m.group(1)]
#         return (-1 if n <= 1 else n - 2), True
#
#     if re.search(r'\b(?:at the (?:beginning|start)|as the first slide|to the front)\b', t):
#         return -1, True
#     if re.search(r'\b(?:at the end|as the last slide|to the end)\b', t):
#         return None, True
#
#     # A bare "slide N" mentioned with add/insert/create phrasing → position N.
#     m = re.search(r'\bslide\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
#     if m and re.search(r'\b(?:add|insert|create|make|new)\b', t):
#         n = int(m.group(1))
#         return (-1 if n <= 1 else n - 2), True
#
#     return None, False
#
#
# _ADD_SLIDE_FILLER = {
#     "add", "insert", "create", "make", "put", "a", "an", "the", "new", "slide", "slides",
#     "please", "can", "you", "could", "would", "for", "me", "to", "us", "we", "want",
#     "need", "like", "and", "one", "here", "in", "presentation", "deck",
# }
#
# def _has_real_topic(text: str) -> bool:
#     """False when the instruction is just the trigger phrase itself ("add a
#     slide", "create a new slide") with no actual subject — the case that
#     used to get sent straight to content generation, which doesn't fail on
#     a topic-less prompt, it just invents plausible-sounding filler."""
#     words = re.findall(r"[a-zA-Z']+", text.lower())
#     meaningful = [w for w in words if w not in _ADD_SLIDE_FILLER and not w.isdigit()]
#     return len(meaningful) >= 2
#
#
# async def ppt_navigate(args: dict, session_id: str) -> dict:
#     direction = args.get("direction", "next")
#     from backend.queues.bus import bus
#     from backend.api.ppt import _slide_store, _latest_upload_sid, _current_slide
#
#     slides = _slide_store.get(session_id) or _slide_store.get(_latest_upload_sid, [])
#     effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
#     total = len(slides)
#     current = _current_slide.get(effective_sid, 0)
#
#     if direction == "next":
#         if total > 0 and current >= total - 1:
#             return {"spoken_reply": f"You've reached the last slide — slide {total} of {total}. That's the end of the presentation."}
#         _current_slide[effective_sid] = min(current + 1, max(total - 1, 0))
#     elif direction == "prev":
#         if current <= 0:
#             return {"spoken_reply": "You're already on the first slide."}
#         _current_slide[effective_sid] = max(current - 1, 0)
#     elif direction == "first":
#         _current_slide[effective_sid] = 0
#     elif direction == "last":
#         _current_slide[effective_sid] = max(total - 1, 0)
#
#     new_idx = _current_slide.get(effective_sid, 0)
#     await bus.emit_event("ppt_command", {"action": direction}, session_id)
#     return {"status": "ok", "direction": direction, "index": new_idx, "spoken_reply": ""}
#
#
# async def ppt_jump_to_title(args: dict, session_id: str) -> dict:
#     query        = args.get("query", "")
#     slide_number = args.get("slide_number")   # already 0-indexed if from keyword fallback
#     from backend.queues.bus import bus
#
#     from backend.api.ppt import _slide_store, _latest_upload_sid, _current_slide
#     import re
#
#     effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
#     slides = _slide_store.get(effective_sid, [])
#
#     # If explicit slide number given, use directly
#     if slide_number is not None:
#         idx = int(slide_number)
#         _current_slide[effective_sid] = idx
#         await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
#         return {"status": "ok", "index": idx, "title": f"Slide {idx+1}", "spoken_reply": ""}
#
#     q = query.lower()
#
#     # Numeric match in query — only allow within actual deck bounds
#     m = re.search(r'\b(\d+)\b', q)
#     if m:
#         idx = int(m.group(1)) - 1
#         if 0 <= idx < len(slides):
#             _current_slide[effective_sid] = idx
#             await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
#             return {"status": "ok", "index": idx, "spoken_reply": ""}
#
#     # Title fuzzy match
#     best_idx, best_score = 0, 0
#     for s in slides:
#         score = sum(1 for w in q.split() if w in s.get("title","").lower())
#         if score > best_score:
#             best_score = score; best_idx = s["index"]
#
#     _current_slide[effective_sid] = best_idx
#     await bus.emit_event("ppt_command", {"action": "goto", "index": best_idx}, session_id)
#     return {"status": "ok", "index": best_idx, "spoken_reply": ""}
#
#
# async def ppt_delete_slide(args: dict, session_id: str) -> dict:
#     from backend.queues.bus import bus
#     slide_number = args.get("slide_number")
#     if slide_number is not None:
#         # Navigate to the target slide first so the confirm modal shows the right one
#         await bus.emit_event("ppt_command", {"action": "goto", "index": int(slide_number)}, session_id)
#     await bus.emit_event("ppt_command", {"action": "delete"}, session_id)
#     label = f"slide {int(slide_number) + 1}" if slide_number is not None else "this slide"
#     return {"status": "ok", "spoken_reply": f"Please confirm the deletion of {label} in the popup."}
#
#
# async def ppt_summarize(args: dict, session_id: str) -> dict:
#     from backend.api.ppt import _slide_store, _latest_upload_sid
#     slides = _slide_store.get(session_id) or _slide_store.get(_latest_upload_sid, [])
#     if not slides:
#         return {"spoken_reply": "No presentation is loaded yet. Please upload a PowerPoint file first."}
#
#     lines = []
#     for s in slides:
#         title = s.get("title", f"Slide {s['index']+1}")
#         notes = s.get("notes", "")
#         lines.append(f"Slide {s['index']+1}: {title}" + (f" — {notes}" if notes else ""))
#
#     content = "\n".join(lines)
#     summary = await _summarize_ollama(content)
#     return {"spoken_reply": summary}
#
#
# async def ppt_last_action(args: dict, session_id: str) -> dict:
#     from backend.core.session_state import get_state
#     state = get_state(session_id)
#     last = state.last_ppt_action or {}
#     if not last:
#         return {"spoken_reply": "I haven't changed the deck yet in this session."}
#     slide_no = last.get("slide_number")
#     instruction = last.get("instruction", "your last request")
#     changes = last.get("changes") or []
#     if changes:
#         changed_text = ", ".join(changes)
#         return {"spoken_reply": f"Yes. I updated slide {slide_no}: {changed_text}, based on your request to {instruction}."}
#     return {"spoken_reply": f"I tried to update slide {slide_no} based on your request to {instruction}, but I don't see a content change."}
#
#
# async def _summarize_ollama(content: str) -> str:
#     def _call() -> str:
#         import ollama
#         from backend.core.config import settings
#         resp = ollama.chat(
#             model=settings.OLLAMA_MODEL,
#             messages=[
#                 {"role": "system", "content":
#                     "You are a helpful voice assistant. Summarize the presentation in 4-6 natural spoken sentences. "
#                     "Mention the main topics and key points. No markdown, no bullet points — plain conversational speech only."},
#                 {"role": "user", "content": f"Summarize this presentation:\n\n{content[:3000]}"},
#             ],
#             think=False,
#             options={"num_predict": 220},
#             stream=False,
#         )
#         if isinstance(resp, dict):
#             return resp["message"]["content"].strip()
#         return resp.message.content.strip()
#
#     try:
#         return await asyncio.to_thread(_call)
#     except Exception as e:
#         logger.error(f"ppt_summarize ollama error: {e}")
#         return "I wasn't able to summarize the presentation right now. Please try again."
#
#
# async def _apply_edit_and_background_refresh(effective_sid: str, idx: int, session_id: str, writer) -> bool:
#     """
#     Writes an edit immediately (fast — well under a second) so the caller
#     can speak a confirmation right away, then regenerates thumbnails
#     (LibreOffice re-rendering the whole deck — several seconds) as a
#     background task, emitting the same "refresh" ppt_command once that
#     finishes instead of making the voice reply wait for it too. The
#     thumbnail visually catches up a couple seconds after the spoken
#     confirmation — the frontend already refetches on that event regardless
#     of when it fires.
#     Returns True on success, False if the session/slide index was invalid.
#     """
#     from backend.api.ppt import apply_slide_edit_fast_async, refresh_slide_thumbnails_async
#     from backend.queues.bus import bus
#
#     ok = await apply_slide_edit_fast_async(effective_sid, idx, writer)
#     if not ok:
#         return False
#
#     async def _background_refresh():
#         await refresh_slide_thumbnails_async(effective_sid)
#         await bus.emit_event("ppt_command", {"action": "refresh", "session_id": effective_sid}, session_id)
#
#     asyncio.create_task(_background_refresh())
#     return True
#
#
# async def ppt_edit_slide(args: dict, session_id: str) -> dict:
#     instruction = args.get("instruction", "")
#     if not instruction:
#         return {"spoken_reply": "What would you like me to change?"}
#
#     from backend.api.ppt import _slide_store, _latest_upload_sid, _current_slide
#     from backend.queues.bus import bus
#     import re
#
#     effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
#     slides = _slide_store.get(effective_sid, [])
#     if not slides:
#         return {"spoken_reply": "No presentation is loaded."}
#
#     slide_number = args.get("slide_number")
#     idx = int(slide_number) if slide_number is not None else _current_slide.get(effective_sid, 0)
#     if idx < 0 or idx >= len(slides):
#         return {"spoken_reply": "I'm not sure which slide to edit."}
#     _current_slide[effective_sid] = idx
#     await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
#
#     slide = slides[idx]
#
#     # Slides generated from the GD template have a known "kind" (team,
#     # table, comparison, ...) whose content doesn't fit the generic
#     # title/numbered-bullets model at all — route those through the
#     # kind-aware editor instead. Uploaded/legacy decks have no kind and
#     # fall straight through to the unchanged generic path below.
#     kind = slide.get("kind")
#     if kind:
#         return await _edit_kind_aware_slide(effective_sid, idx, kind, instruction, session_id, slide.get("source"))
#
#     # Extract bullets
#     old_bullets = []
#     for sh in slide.get("shapes", []):
#         for line in sh.get("text", "").split("\n"):
#             m = re.match(r'^\d+\.\s{1,3}(.+)', line)
#             if m:
#                 old_bullets.append(m.group(1).strip())
#
#     old_title = slide.get("title", "")
#     old_notes = slide.get("notes", "")
#
#     # Every other text-bearing shape that ISN'T the title or a numbered bullet
#     # line (e.g. a subtitle, caption, or footer). Keyed by stable shape_id so
#     # the editor can target one exact shape without touching anything else on
#     # the slide — this is what lets the user edit "anything on the slide",
#     # not just the title/bullets/notes triplet.
#     other_shapes: dict[str, str] = {}
#     for sh in slide.get("shapes", []):
#         sid = sh.get("shape_id")
#         text = sh.get("text", "").strip()
#         if sid is None or not text:
#             continue
#         if text == old_title:
#             continue
#         if all(re.match(r'^\d+\.\s{1,3}', line) for line in text.split("\n") if line.strip()):
#             continue  # this is the bullet body shape, already covered above
#         other_shapes[str(sid)] = text
#
#     # Common voice correction: "change the title of slide 3 from X to Y".
#     # Handle it deterministically so ASR/correction commands don't depend on
#     # the editor LLM preserving the exact replacement text.
#     title_to = re.search(r'\btitle\b.*?\bfrom\b.+?\bto\b\s+(.+)$', instruction, re.IGNORECASE)
#     title_direct = re.search(r'\b(?:change|update|set|rename)\b.*?\btitle\b.*?\bto\b\s+(.+)$', instruction, re.IGNORECASE)
#     replacement_title = (title_to or title_direct)
#     if replacement_title:
#         new_title = replacement_title.group(1).strip().strip('".,!?')
#         if new_title:
#             try:
#                 from backend.api.ppt import _patch_slide
#                 # other_edits intentionally omitted (defaults to None): a
#                 # title-only instruction must never touch bullets, notes, or
#                 # any other shape on the slide.
#                 def _writer(slide, _t=new_title, _b=old_bullets, _n=old_notes):
#                     _patch_slide(slide, idx, _t, _b, _n, None)
#                 await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
#                 from backend.core.session_state import get_state
#                 state = get_state(session_id)
#                 state.last_ppt_action = {
#                     "tool": "ppt_edit_slide",
#                     "slide_number": idx + 1,
#                     "instruction": instruction,
#                     "changes": [f"title from '{old_title}' to '{new_title}'"],
#                 }
#                 return {"spoken_reply": f"I changed the title on slide {idx + 1} to {new_title}."}
#             except Exception as e:
#                 logger.error(f"ppt_edit_slide title shortcut error: {e}")
#                 return {"spoken_reply": "I had trouble editing the slide right now."}
#
#     # LLM call
#     def _call() -> str:
#         import ollama
#         from backend.core.config import settings
#         import json
#         other_shapes_desc = (
#             "\n".join(f'- shape_id {sid}: "{text}"' for sid, text in other_shapes.items())
#             if other_shapes else "(none)"
#         )
#         prompt = f"""
# You are an expert presentation editor. The user wants to edit a slide.
# Current Slide:
# Title: {old_title}
# Bullets: {json.dumps(old_bullets)}
# Notes: {old_notes}
# Other text on this slide (subtitle/caption/footer — NOT the title or bullets):
# {other_shapes_desc}
#
# Instruction: {instruction}
#
# Rules:
# - Only change what the instruction actually asks for. Fields you are not asked
#   to change must be echoed back EXACTLY as given above — do not paraphrase,
#   shorten, or drop unrelated text.
# - If the instruction targets one of the "other text" shapes above (a subtitle,
#   caption, footer, etc.), put its new text in "other_edits" keyed by its
#   shape_id. Do not invent shape_ids that weren't listed. Leave "other_edits"
#   as {{}} if nothing else needs to change.
#
# Output ONLY valid JSON matching this schema:
# {{
#   "title": "New Title",
#   "bullets": ["New Bullet 1", "New Bullet 2"],
#   "notes": "New Speaker Notes",
#   "other_edits": {{"<shape_id>": "new text for that shape"}}
# }}
# """
#         resp = ollama.chat(
#             model=settings.OLLAMA_MODEL,
#             messages=[{"role": "user", "content": prompt}],
#             # think=False is required alongside format="json" — without it, a
#             # reasoning-capable model (Qwen3, etc.) tries to emit chain-of-thought
#             # before its JSON, which conflicts with grammar-constrained JSON
#             # decoding and can come back as empty content (see _summarize_ollama
#             # above, and _generate_template_content_sync in
#             # services/ppt_template_builder.py, for the same fix).
#             options={"num_predict": 400},
#             format="json",
#             think=False,
#             stream=False,
#         )
#         if isinstance(resp, dict):
#             return resp["message"]["content"].strip()
#         return resp.message.content.strip()
#
#     try:
#         raw = await asyncio.to_thread(_call)
#         import json
#         if not raw:
#             # Same empty-content failure mode this fix targets — retry once
#             # before giving up, since Ollama occasionally still returns a
#             # blank message.strip() on a cold model load.
#             raw = await asyncio.to_thread(_call)
#         if not raw:
#             logger.error("ppt_edit_slide error: Ollama returned an empty response after retry")
#             return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}
#         try:
#             data = json.loads(raw)
#         except json.JSONDecodeError:
#             logger.error(f"ppt_edit_slide error: invalid JSON from Ollama: {raw[:300]!r}")
#             return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}
#         new_title = data.get("title", old_title)
#         new_bullets = data.get("bullets", old_bullets)
#         new_notes = data.get("notes", old_notes)
#
#         # Only accept other_edits for shape_ids we actually listed to the
#         # model — guards against a hallucinated id landing on the wrong shape.
#         raw_other_edits = data.get("other_edits") or {}
#         other_edits = {
#             sid: text for sid, text in raw_other_edits.items()
#             if sid in other_shapes and isinstance(text, str) and text.strip()
#         }
#
#         changes = []
#         if new_title != old_title:
#             changes.append(f"title from '{old_title}' to '{new_title}'")
#         if new_bullets != old_bullets:
#             changes.append("bullet content")
#         if new_notes != old_notes:
#             changes.append("speaker notes")
#         for sid, text in other_edits.items():
#             if text != other_shapes.get(sid):
#                 changes.append("other slide text")
#                 break
#
#         from backend.api.ppt import _patch_slide
#         def _writer(slide, _t=new_title, _b=new_bullets, _n=new_notes, _oe=other_edits or None):
#             _patch_slide(slide, idx, _t, _b, _n, _oe)
#         await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
#         from backend.core.session_state import get_state
#         state = get_state(session_id)
#         state.last_ppt_action = {
#             "tool": "ppt_edit_slide",
#             "slide_number": idx + 1,
#             "instruction": instruction,
#             "changes": changes,
#         }
#         return {"spoken_reply": f"I've updated slide {idx + 1}."}
#     except Exception as e:
#         logger.error(f"ppt_edit_slide error: {e}")
#         return {"spoken_reply": "I had trouble editing the slide right now."}
#
#
# async def _edit_kind_aware_slide(effective_sid: str, idx: int, kind: str, instruction: str, session_id: str, source: str | None = None) -> dict:
#     """
#     Voice-edit path for slides cloned from the GD template (team, table,
#     comparison, ...) — the generic title/numbered-bullets model above
#     doesn't fit their structure at all. Same "send current state as JSON,
#     ask the model for edited JSON back" pattern as the generic path, just
#     with each kind's own field schema instead of title/bullets/notes.
#
#     `source` identifies exactly which candidate slide (of possibly several
#     for this kind — see _KIND_SOURCES) this one was originally cloned from;
#     different candidates don't share shape_ids, so it's needed to compute
#     the right slot map, not just the kind name.
#     """
#     from backend.services.ppt_template_builder import extract_slide_data, populate_slide_data, _parse_source
#     import json
#
#     parsed_source = _parse_source(source)
#
#     def _read_current() -> dict:
#         from pptx import Presentation
#         prs = Presentation(f"data/ppt/{effective_sid}.pptx")
#         return extract_slide_data(prs.slides[idx], kind, parsed_source)
#
#     try:
#         current_data = await asyncio.to_thread(_read_current)
#     except Exception as e:
#         logger.error(f"kind-aware edit: failed to read current slide data: {e}")
#         return {"spoken_reply": "I had trouble reading that slide right now."}
#
#     def _call() -> str:
#         import ollama
#         from backend.core.config import settings
#         prompt = f"""
# You are an expert presentation editor. The user wants to edit a "{kind}" slide.
#
# Current content (JSON): {json.dumps(current_data)}
#
# Instruction: {instruction}
#
# Rules:
# - Return the SAME JSON shape as "Current content" above, with the requested
#   change applied.
# - Only change what the instruction actually asks for — every other field
#   must be echoed back EXACTLY as given, do not paraphrase or drop anything.
# - Output ONLY valid JSON — no markdown, no explanation, nothing else.
# """
#         resp = ollama.chat(
#             model=settings.OLLAMA_MODEL,
#             messages=[{"role": "user", "content": prompt}],
#             options={"num_predict": 500},
#             format="json",
#             think=False,  # see ppt_edit_slide above for why this must pair with format="json"
#             stream=False,
#         )
#         if isinstance(resp, dict):
#             return resp["message"]["content"].strip()
#         return resp.message.content.strip()
#
#     try:
#         raw = await asyncio.to_thread(_call)
#         if not raw:
#             raw = await asyncio.to_thread(_call)  # cold-model-load retry, same as the generic path
#         if not raw:
#             logger.error("kind-aware edit error: Ollama returned an empty response after retry")
#             return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}
#         try:
#             new_data = json.loads(raw)
#         except json.JSONDecodeError:
#             logger.error(f"kind-aware edit error: invalid JSON from Ollama: {raw[:300]!r}")
#             return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}
#
#         def _writer(slide):
#             populate_slide_data(slide, kind, new_data, parsed_source)
#
#         await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
#         from backend.core.session_state import get_state
#         state = get_state(session_id)
#         state.last_ppt_action = {
#             "tool": "ppt_edit_slide",
#             "slide_number": idx + 1,
#             "instruction": instruction,
#             "changes": ["slide content"],
#         }
#         return {"spoken_reply": f"I've updated slide {idx + 1}."}
#     except Exception as e:
#         logger.error(f"kind-aware ppt_edit_slide error: {e}")
#         return {"spoken_reply": "I had trouble editing the slide right now."}
#
#
# def _slide_bullets_for_notes(slide: dict) -> list:
#     import re
#     bullets = []
#     for sh in slide.get("shapes", []):
#         for line in sh.get("text", "").split("\n"):
#             m = re.match(r'^\d+\.\s{1,3}(.+)', line)
#             if m:
#                 bullets.append(m.group(1).strip())
#     return bullets
#
#
# def _generate_notes_text_sync(title: str, bullets: list) -> str:
#     import ollama, json
#     from backend.core.config import settings
#     prompt = f"""
# Write 3-4 natural conversational sentences of speaker notes for this slide.
# Title: {title}
# Bullets: {json.dumps(bullets)}
#
# Output the speaker notes plainly without any markdown, prefix, or JSON formatting.
# """
#     resp = ollama.chat(
#         model=settings.OLLAMA_MODEL,
#         messages=[{"role": "user", "content": prompt}],
#         options={"num_predict": 250},
#         stream=False,
#     )
#     if isinstance(resp, dict):
#         return resp["message"]["content"].strip()
#     return resp.message.content.strip()
#
#
# async def ppt_generate_notes(args: dict, session_id: str) -> dict:
#     from backend.api.ppt import (
#         _slide_store, _latest_upload_sid, _current_slide,
#         apply_notes_batch_async,
#     )
#     from backend.queues.bus import bus
#
#     effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
#     slides = _slide_store.get(effective_sid, [])
#     if not slides:
#         return {"spoken_reply": "No presentation is loaded."}
#
#     if args.get("all"):
#         await bus.emit_event("ppt_command", {"action": "goto", "index": 0}, session_id)
#         notes_by_index: dict[int, str] = {}
#         failures = 0
#         for i, slide in enumerate(slides):
#             try:
#                 notes_by_index[i] = await asyncio.to_thread(
#                     _generate_notes_text_sync, slide.get("title", ""), _slide_bullets_for_notes(slide),
#                 )
#             except Exception as e:
#                 failures += 1
#                 logger.error(f"ppt_generate_notes (all) slide {i} error: {e}")
#         if not notes_by_index:
#             return {"spoken_reply": "I had trouble generating notes right now."}
#         await apply_notes_batch_async(effective_sid, notes_by_index)
#         await bus.emit_event("ppt_command", {"action": "refresh", "session_id": effective_sid}, session_id)
#         from backend.core.session_state import get_state
#         state = get_state(session_id)
#         state.last_ppt_action = {
#             "tool": "ppt_generate_notes",
#             "slide_number": None,
#             "instruction": "generate speaker notes for all slides",
#             "changes": [f"speaker notes for {len(notes_by_index)} of {len(slides)} slides"],
#         }
#         if failures:
#             return {"spoken_reply": f"I've generated speaker notes for {len(notes_by_index)} of {len(slides)} slides — {failures} failed."}
#         return {"spoken_reply": f"I've generated speaker notes for all {len(slides)} slides."}
#
#     slide_number = args.get("slide_number")
#     idx = int(slide_number) if slide_number is not None else _current_slide.get(effective_sid, 0)
#     if idx < 0 or idx >= len(slides):
#         return {"spoken_reply": "I'm not sure which slide to write notes for."}
#     _current_slide[effective_sid] = idx
#     await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
#     slide = slides[idx]
#     old_title = slide.get("title", "")
#     old_bullets = _slide_bullets_for_notes(slide)
#
#     try:
#         notes = await asyncio.to_thread(_generate_notes_text_sync, old_title, old_bullets)
#         # Apply notes, keep title and bullets the same
#         from backend.api.ppt import _patch_slide
#         def _writer(slide, _t=old_title, _b=old_bullets, _n=notes):
#             _patch_slide(slide, idx, _t, _b, _n, None)
#         await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
#         from backend.core.session_state import get_state
#         state = get_state(session_id)
#         state.last_ppt_action = {
#             "tool": "ppt_generate_notes",
#             "slide_number": idx + 1,
#             "instruction": "generate speaker notes",
#             "changes": ["speaker notes"],
#         }
#         return {"spoken_reply": f"I've generated new speaker notes for slide {idx + 1}."}
#     except Exception as e:
#         logger.error(f"ppt_generate_notes error: {e}")
#         return {"spoken_reply": "I had trouble generating notes right now."}
#
#
# async def ppt_add_slide(args: dict, session_id: str) -> dict:
#     """Insert a new slide into the loaded presentation, content generated
#     from a spoken description ("add a slide about our Q4 roadmap [after
#     slide 3]"). If the instruction has no real topic ("add a slide" with
#     nothing else), asks what it should be about instead of sending that
#     bare trigger phrase to content generation — which doesn't fail on a
#     topic-less prompt, it just invents plausible-sounding filler that reads
#     as if the system "hallucinated" or echoed something already in the
#     deck. The clarifying answer is picked up on the next turn via
#     state.pending_add_slide (see services/front_llm.py's classify())."""
#     from backend.api.ppt import _slide_store, _latest_upload_sid
#     from backend.services.ppt_template_builder import generate_single_slide_content
#     from backend.core.session_state import get_state
#
#     effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
#     slides = _slide_store.get(effective_sid, [])
#     if not slides:
#         return {"spoken_reply": "No presentation is loaded yet. Please upload or create one first."}
#
#     state = get_state(session_id)
#     pending = state.pending_add_slide
#     instruction = args.get("instruction", "").strip()
#
#     insert_after, position_specified = _parse_insert_position(instruction)
#     # Whether the user ever named a position — this turn, or in the pending
#     # clarification round. If they did, we honour it exactly; if they never
#     # did, PILOT picks the best spot itself (from the slide titles) so the
#     # deck keeps a coherent flow, and tells the user where it landed.
#     if pending is not None:
#         if position_specified:
#             pass  # user gave a position on the clarification turn — use it
#         elif pending.get("position_specified"):
#             insert_after = pending.get("insert_after")
#             position_specified = True
#     user_chose_position = position_specified
#
#     if not _has_real_topic(instruction):
#         if pending is not None:
#             # Already asked once this round and still got nothing to go on
#             # — don't loop forever asking the same question.
#             state.pending_add_slide = None
#             return {"spoken_reply": "I still didn't catch what the slide should be about, so I'll leave it for now — just ask again whenever you're ready."}
#         # Remember whether a position was already stated so we don't re-ask or
#         # override it after we get the topic on the next turn.
#         state.pending_add_slide = {"insert_after": insert_after, "position_specified": position_specified}
#         where = "" if not position_specified else (
#             "at the very beginning" if insert_after == -1
#             else "at the end" if insert_after is None
#             else f"right after slide {insert_after + 1}"
#         )
#         return {"spoken_reply": "Sure — what should the new slide be about?" + (f" I'll put it {where}." if where else "")}
#
#     state.pending_add_slide = None
#
#     try:
#         slide_data = await generate_single_slide_content(instruction)
#     except Exception as e:
#         logger.error(f"ppt_add_slide generation error: {e}")
#         slide_data = None
#
#     if not slide_data:
#         # Generation failed/timed out — still produce a real GD-template slide,
#         # not a blank one: "text" is a valid template kind whose source slides
#         # carry the Grid Dynamics branding.
#         logger.warning("ppt_add_slide: generation returned nothing — using GD 'text' fallback")
#         slide_data = {"kind": "text", "title": instruction[:60], "paragraphs": [instruction]}
#
#     kind = slide_data.pop("kind", None)
#
#     # GUARANTEE the GD template is used: every added slide must be one of the
#     # known template kinds (each maps to real GD template source slides in
#     # _KIND_SOURCES). If the model returned an unknown/blank kind, coerce to
#     # "text" rather than letting add_slide_to_deck raise and — worse — ever
#     # emit an off-brand slide.
#     from backend.services.ppt_template_builder import _KIND_SOURCES
#     if kind not in _KIND_SOURCES:
#         logger.warning(f"ppt_add_slide: invalid kind {kind!r} — coercing to GD 'text'")
#         title = slide_data.get("title") or instruction[:60]
#         slide_data = {"title": title, "paragraphs": slide_data.get("paragraphs") or [instruction]}
#         kind = "text"
#     new_title = slide_data.get("title") or instruction[:60]
#
#     # ── Auto-position: the user didn't say where, so read the existing slide
#     # titles and let PILOT choose the spot that best preserves the narrative
#     # flow, then announce it. Falls back to appending at the end. ──
#     placement_note = ""
#     if not user_chose_position:
#         from backend.services.ppt_template_builder import pick_insert_position
#         titles = [(s.get("title") or "") for s in slides]
#         topic = new_title if kind not in ("agenda",) else (instruction or new_title)
#         picked = await pick_insert_position(titles, topic)
#         if picked is not None:
#             insert_after, reason = picked
#             if insert_after == -1:
#                 placement_note = " I placed it at the start"
#             elif insert_after is None:
#                 placement_note = " I placed it at the end"
#             else:
#                 placement_note = f" I placed it after slide {insert_after + 1}"
#             placement_note += f" — {reason}." if reason else "."
#         else:
#             insert_after = None  # append at end
#             placement_note = " I added it at the end."
#
#     try:
#         ok = await _apply_add_and_background_refresh(effective_sid, kind, slide_data, session_id, insert_after)
#         if not ok:
#             return {"spoken_reply": "I couldn't add a slide right now."}
#         state.last_ppt_action = {
#             "tool": "ppt_add_slide",
#             "slide_number": len(slides) + 1,
#             "instruction": instruction or "add a new slide",
#             "changes": [f"new {kind} slide"],
#         }
#         return {"spoken_reply": f"I've added a new slide — {new_title}.{placement_note}"}
#     except Exception as e:
#         logger.error(f"ppt_add_slide error: {e}")
#         return {"spoken_reply": "I had trouble adding that slide."}
#
#
# async def _apply_add_and_background_refresh(effective_sid: str, kind: str, data: dict, session_id: str, insert_after: int | None = None) -> bool:
#     """Same fast-reply-then-background-refresh pattern as
#     _apply_edit_and_background_refresh, for the add-slide path: writes the
#     new slide immediately, then regenerates thumbnails (LibreOffice, whole
#     deck) in the background and re-emits ppt_command so the viewer catches
#     up a moment after the spoken confirmation instead of before it."""
#     from backend.api.ppt import add_slide_fast_async, refresh_slide_thumbnails_async
#     from backend.queues.bus import bus
#
#     new_index = await add_slide_fast_async(effective_sid, kind, data, insert_after)
#     if new_index is None:
#         return False
#
#     async def _background_refresh():
#         await refresh_slide_thumbnails_async(effective_sid)
#         await bus.emit_event("ppt_command", {"action": "goto", "index": new_index}, session_id)
#         await bus.emit_event("ppt_command", {"action": "refresh", "session_id": effective_sid}, session_id)
#
#     asyncio.create_task(_background_refresh())
#     return True
#
#
# def _parse_reorder(text: str, total: int) -> tuple[int, int] | None:
#     """Parse 'move slide X after/before slide Y', 'move slide X to position Y',
#     'move slide X to the front/end', ordinal forms. Returns (from_index,
#     to_index) both 0-based, or None if it can't find a clear source+target."""
#     t = text.lower()
#
#     def _num(word_or_digit: str) -> int | None:
#         if word_or_digit.isdigit():
#             return int(word_or_digit)
#         return _ORDINALS.get(word_or_digit)
#
#     # Source slide: "move slide 2" / "move the second slide" / "move slides 1 and 2"
#     src = None
#     m = re.search(r'\bslide[s]?\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
#     if m:
#         src = int(m.group(1))
#     else:
#         m = re.search(r'\b(' + "|".join(_ORDINALS) + r')\s+slide\b', t)
#         if m:
#             src = _ORDINALS[m.group(1)]
#     if src is None or not (1 <= src <= total):
#         return None
#     from_index = src - 1
#
#     # Target: after/before slide Y, to position Y, to front/end.
#     # _move_slide removes the source, THEN inserts. So compute the target
#     # against the POST-removal list: the anchor slide Y sits at index (y-1),
#     # shifted down by one if the source was before it.
#     def _anchor_index(y: int) -> int:
#         idx = y - 1
#         return idx - 1 if idx > from_index else idx
#
#     m = re.search(r'\bafter\s+slide\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
#     if m:
#         to = _anchor_index(int(m.group(1))) + 1  # land right AFTER the anchor
#         return from_index, max(0, min(to, total - 1))
#     m = re.search(r'\bbefore\s+slide\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
#     if m:
#         to = _anchor_index(int(m.group(1)))       # land AT the anchor's slot (pushes it down)
#         return from_index, max(0, min(to, total - 1))
#     m = re.search(r'\b(?:to|at|into?|position)\s+(?:the\s+)?(?:slide|position)?\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
#     if m:
#         return from_index, max(0, min(int(m.group(1)) - 1, total - 1))
#     if re.search(r'\b(?:to the (?:front|beginning|start)|as the first)\b', t):
#         return from_index, 0
#     if re.search(r'\b(?:to the end|as the last)\b', t):
#         return from_index, total - 1
#     return None
#
#
# async def ppt_reorder_slide(args: dict, session_id: str) -> dict:
#     """Reorder a slide within the loaded deck ('move slide 1 after slide 3',
#     'move slide 2 to position 4'). This is the real reorder capability — before
#     it existed, such requests fell through to a generic answer that invented
#     manual PowerPoint drag-and-drop steps."""
#     from backend.api.ppt import _slide_store, _latest_upload_sid, reorder_slide_fast_async, refresh_slide_thumbnails_async
#     from backend.queues.bus import bus
#
#     effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
#     slides = _slide_store.get(effective_sid, [])
#     total = len(slides)
#     if total < 2:
#         return {"spoken_reply": "There aren't enough slides to reorder yet."}
#
#     instruction = args.get("instruction", "")
#     parsed = _parse_reorder(instruction, total)
#     if parsed is None:
#         return {"spoken_reply": "Tell me which slide to move and where — for example, “move slide 1 after slide 3.”"}
#     from_index, to_index = parsed
#     if from_index == to_index:
#         return {"spoken_reply": f"Slide {from_index + 1} is already in that position."}
#
#     try:
#         final_index = await reorder_slide_fast_async(effective_sid, from_index, to_index)
#         if final_index is None:
#             return {"spoken_reply": "No presentation is loaded to reorder."}
#
#         async def _bg():
#             await refresh_slide_thumbnails_async(effective_sid)
#             await bus.emit_event("ppt_command", {"action": "goto", "index": final_index}, session_id)
#             await bus.emit_event("ppt_command", {"action": "refresh", "session_id": effective_sid}, session_id)
#         asyncio.create_task(_bg())
#
#         # NOTE: upstream (c3a22cfc3) is missing this import at the call site
#         # below — every other tool in this file imports get_state locally
#         # before using it; without it this would raise NameError the first
#         # time a reorder actually succeeds.
#         from backend.core.session_state import get_state
#         get_state(session_id).last_ppt_action = {
#             "tool": "ppt_reorder_slide", "slide_number": final_index + 1,
#             "instruction": instruction, "changes": [f"moved slide {from_index + 1} → position {final_index + 1}"],
#         }
#         return {"spoken_reply": f"Moved slide {from_index + 1} to position {final_index + 1}."}
#     except ValueError as e:
#         return {"spoken_reply": str(e)}
#     except Exception as e:
#         logger.error(f"ppt_reorder_slide error: {e}")
#         return {"spoken_reply": "I had trouble reordering that slide."}
# ============================================================================
# ACTIVE: ppt_copilot.py transferred from pilot-voice-agent-backend-ppt_update
# (newer snapshot, supersedes the feature/ppt-copilot @ 869dd08d9 version
# above). Import paths adapted to PILOT's backend.-prefixed convention; the
# missing get_state import bugfix in ppt_reorder_slide (see prior version's
# comment) has been re-applied.
# ============================================================================
"""PPT tools — navigate + jump to slide by number or title + summarize."""
import asyncio, logging, re
logger = logging.getLogger("pilot.tools.ppt")


# ── Add-slide helpers — parse "where" out of the instruction, and detect
# when "what" was never actually said (so we ask instead of hallucinating
# generic filler content that only coincidentally resembles a real slide) ──

_ORDINALS = {
    "first": 1, "second": 2, "third": 3, "fourth": 4, "fifth": 5,
    "sixth": 6, "seventh": 7, "eighth": 8, "ninth": 9, "tenth": 10,
}

def _parse_insert_position(text: str) -> tuple[int | None, bool]:
    """Returns (insert_after, was_specified). insert_after is 0-indexed —
    the new slide lands right after prs.slides[insert_after]. -1 means "at the
    very beginning" (before slide 1), None means "append at the end" (also the
    default when nothing was said)."""
    t = text.lower()

    # "after slide N" → land right after slide N (0-indexed N-1).
    m = re.search(r'\bafter\s+slide\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
    if m:
        return int(m.group(1)) - 1, True

    # "before slide N" → before slide N == after slide N-1 (0-indexed N-2).
    m = re.search(r'\bbefore\s+slide\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
    if m:
        return int(m.group(1)) - 2, True

    # "on/at slide N", "at position N", "as slide N", "make it slide N",
    # "in position N", or a bare "slide N" — all mean "the new slide should
    # BECOME slide N", i.e. occupy position N. That's after slide N-1 (0-indexed
    # N-2), so slide N-1 stays before it and the old slide N shifts down.
    m = re.search(
        r'\b(?:on|at|as|in(?:to)?|position|make\s+it)\s+(?:the\s+)?'
        r'(?:slide|position)?\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
    if m:
        n = int(m.group(1))
        return (-1 if n <= 1 else n - 2), True

    # Ordinal word forms: "as the third slide", "make it the second slide".
    m = re.search(r'\b(?:as|make\s+it|at|position)\s+(?:the\s+)?(' + "|".join(_ORDINALS) + r')\s+slide\b', t)
    if m:
        n = _ORDINALS[m.group(1)]
        return (-1 if n <= 1 else n - 2), True

    if re.search(r'\b(?:at the (?:beginning|start)|as the first slide|to the front)\b', t):
        return -1, True
    if re.search(r'\b(?:at the end|as the last slide|to the end)\b', t):
        return None, True

    # A bare "slide N" mentioned with add/insert/create phrasing → position N.
    m = re.search(r'\bslide\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
    if m and re.search(r'\b(?:add|insert|create|make|new)\b', t):
        n = int(m.group(1))
        return (-1 if n <= 1 else n - 2), True

    return None, False


_ADD_SLIDE_FILLER = {
    "add", "insert", "create", "make", "put", "a", "an", "the", "new", "slide", "slides",
    "please", "can", "you", "could", "would", "for", "me", "to", "us", "we", "want",
    "need", "like", "and", "one", "here", "in", "presentation", "deck",
}

def _has_real_topic(text: str) -> bool:
    """False when the instruction is just the trigger phrase itself ("add a
    slide", "create a new slide") with no actual subject — the case that
    used to get sent straight to content generation, which doesn't fail on
    a topic-less prompt, it just invents plausible-sounding filler."""
    words = re.findall(r"[a-zA-Z']+", text.lower())
    meaningful = [w for w in words if w not in _ADD_SLIDE_FILLER and not w.isdigit()]
    return len(meaningful) >= 2


async def ppt_navigate(args: dict, session_id: str) -> dict:
    direction = args.get("direction", "next")
    from backend.queues.bus import bus
    from backend.api.ppt import _slide_store, _latest_upload_sid, _current_slide

    slides = _slide_store.get(session_id) or _slide_store.get(_latest_upload_sid, [])
    effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
    total = len(slides)
    current = _current_slide.get(effective_sid, 0)

    if direction == "next":
        if total > 0 and current >= total - 1:
            return {"spoken_reply": f"You've reached the last slide — slide {total} of {total}. That's the end of the presentation."}
        _current_slide[effective_sid] = min(current + 1, max(total - 1, 0))
    elif direction == "prev":
        if current <= 0:
            return {"spoken_reply": "You're already on the first slide."}
        _current_slide[effective_sid] = max(current - 1, 0)
    elif direction == "first":
        _current_slide[effective_sid] = 0
    elif direction == "last":
        _current_slide[effective_sid] = max(total - 1, 0)

    new_idx = _current_slide.get(effective_sid, 0)
    await bus.emit_event("ppt_command", {"action": direction}, session_id)
    return {"status": "ok", "direction": direction, "index": new_idx, "spoken_reply": ""}


async def ppt_jump_to_title(args: dict, session_id: str) -> dict:
    query        = args.get("query", "")
    slide_number = args.get("slide_number")   # already 0-indexed if from keyword fallback
    from backend.queues.bus import bus

    from backend.api.ppt import _slide_store, _latest_upload_sid, _current_slide
    import re

    effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
    slides = _slide_store.get(effective_sid, [])

    # If explicit slide number given, use directly
    if slide_number is not None:
        idx = int(slide_number)
        _current_slide[effective_sid] = idx
        await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
        return {"status": "ok", "index": idx, "title": f"Slide {idx+1}", "spoken_reply": ""}

    q = query.lower()

    # Numeric match in query — only allow within actual deck bounds
    m = re.search(r'\b(\d+)\b', q)
    if m:
        idx = int(m.group(1)) - 1
        if 0 <= idx < len(slides):
            _current_slide[effective_sid] = idx
            await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
            return {"status": "ok", "index": idx, "spoken_reply": ""}

    # Title fuzzy match
    best_idx, best_score = 0, 0
    for s in slides:
        score = sum(1 for w in q.split() if w in s.get("title","").lower())
        if score > best_score:
            best_score = score; best_idx = s["index"]

    _current_slide[effective_sid] = best_idx
    await bus.emit_event("ppt_command", {"action": "goto", "index": best_idx}, session_id)
    return {"status": "ok", "index": best_idx, "spoken_reply": ""}


async def ppt_delete_slide(args: dict, session_id: str) -> dict:
    from backend.queues.bus import bus
    slide_number = args.get("slide_number")
    if slide_number is not None:
        # Navigate to the target slide first so the confirm modal shows the right one
        await bus.emit_event("ppt_command", {"action": "goto", "index": int(slide_number)}, session_id)
    await bus.emit_event("ppt_command", {"action": "delete"}, session_id)
    label = f"slide {int(slide_number) + 1}" if slide_number is not None else "this slide"
    return {"status": "ok", "spoken_reply": f"Please confirm the deletion of {label} in the popup."}


async def ppt_summarize(args: dict, session_id: str) -> dict:
    from backend.api.ppt import _slide_store, _latest_upload_sid
    slides = _slide_store.get(session_id) or _slide_store.get(_latest_upload_sid, [])
    if not slides:
        return {"spoken_reply": "No presentation is loaded yet. Please upload a PowerPoint file first."}

    lines = []
    for s in slides:
        title = s.get("title", f"Slide {s['index']+1}")
        notes = s.get("notes", "")
        lines.append(f"Slide {s['index']+1}: {title}" + (f" — {notes}" if notes else ""))

    content = "\n".join(lines)
    summary = await _summarize_ollama(content)
    return {"spoken_reply": summary}


async def ppt_last_action(args: dict, session_id: str) -> dict:
    from backend.core.session_state import get_state
    state = get_state(session_id)
    last = state.last_ppt_action or {}
    if not last:
        return {"spoken_reply": "I haven't changed the deck yet in this session."}
    slide_no = last.get("slide_number")
    instruction = last.get("instruction", "your last request")
    changes = last.get("changes") or []
    if changes:
        changed_text = ", ".join(changes)
        return {"spoken_reply": f"Yes. I updated slide {slide_no}: {changed_text}, based on your request to {instruction}."}
    return {"spoken_reply": f"I tried to update slide {slide_no} based on your request to {instruction}, but I don't see a content change."}


async def _summarize_ollama(content: str) -> str:
    def _call() -> str:
        import ollama
        from backend.core.config import settings
        resp = ollama.chat(
            model=settings.OLLAMA_MODEL,
            messages=[
                {"role": "system", "content":
                    "You are a helpful voice assistant. Summarize the presentation in 4-6 natural spoken sentences. "
                    "Mention the main topics and key points. No markdown, no bullet points — plain conversational speech only."},
                {"role": "user", "content": f"Summarize this presentation:\n\n{content[:3000]}"},
            ],
            think=False,
            options={"num_predict": 220},
            stream=False,
        )
        if isinstance(resp, dict):
            return resp["message"]["content"].strip()
        return resp.message.content.strip()

    try:
        return await asyncio.to_thread(_call)
    except Exception as e:
        logger.error(f"ppt_summarize ollama error: {e}")
        return "I wasn't able to summarize the presentation right now. Please try again."


async def _apply_edit_and_background_refresh(effective_sid: str, idx: int, session_id: str, writer) -> bool:
    """
    Writes an edit immediately (fast — well under a second) so the caller
    can speak a confirmation right away, then regenerates thumbnails
    (LibreOffice re-rendering the whole deck — several seconds) as a
    background task, emitting the same "refresh" ppt_command once that
    finishes instead of making the voice reply wait for it too. The
    thumbnail visually catches up a couple seconds after the spoken
    confirmation — the frontend already refetches on that event regardless
    of when it fires.
    Returns True on success, False if the session/slide index was invalid.
    """
    from backend.api.ppt import apply_slide_edit_fast_async, refresh_slide_thumbnails_async
    from backend.queues.bus import bus

    ok = await apply_slide_edit_fast_async(effective_sid, idx, writer)
    if not ok:
        return False

    async def _background_refresh():
        await refresh_slide_thumbnails_async(effective_sid)
        await bus.emit_event("ppt_command", {"action": "refresh", "session_id": effective_sid}, session_id)

    asyncio.create_task(_background_refresh())
    return True


async def ppt_edit_slide(args: dict, session_id: str) -> dict:
    instruction = args.get("instruction", "")
    if not instruction:
        return {"spoken_reply": "What would you like me to change?"}

    from backend.api.ppt import _slide_store, _latest_upload_sid, _current_slide
    from backend.queues.bus import bus
    import re

    effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
    slides = _slide_store.get(effective_sid, [])
    if not slides:
        return {"spoken_reply": "No presentation is loaded."}
        
    slide_number = args.get("slide_number")
    idx = int(slide_number) if slide_number is not None else _current_slide.get(effective_sid, 0)
    if idx < 0 or idx >= len(slides):
        return {"spoken_reply": "I'm not sure which slide to edit."}
    _current_slide[effective_sid] = idx
    await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)

    slide = slides[idx]

    # Slides generated from the GD template have a known "kind" (team,
    # table, comparison, ...) whose content doesn't fit the generic
    # title/numbered-bullets model at all — route those through the
    # kind-aware editor instead. Uploaded/legacy decks have no kind and
    # fall straight through to the unchanged generic path below.
    kind = slide.get("kind")
    if kind:
        return await _edit_kind_aware_slide(effective_sid, idx, kind, instruction, session_id, slide.get("source"))

    # Extract bullets
    old_bullets = []
    for sh in slide.get("shapes", []):
        for line in sh.get("text", "").split("\n"):
            m = re.match(r'^\d+\.\s{1,3}(.+)', line)
            if m:
                old_bullets.append(m.group(1).strip())
                
    old_title = slide.get("title", "")
    old_notes = slide.get("notes", "")

    # Every other text-bearing shape that ISN'T the title or a numbered bullet
    # line (e.g. a subtitle, caption, or footer). Keyed by stable shape_id so
    # the editor can target one exact shape without touching anything else on
    # the slide — this is what lets the user edit "anything on the slide",
    # not just the title/bullets/notes triplet.
    other_shapes: dict[str, str] = {}
    for sh in slide.get("shapes", []):
        sid = sh.get("shape_id")
        text = sh.get("text", "").strip()
        if sid is None or not text:
            continue
        if text == old_title:
            continue
        if all(re.match(r'^\d+\.\s{1,3}', line) for line in text.split("\n") if line.strip()):
            continue  # this is the bullet body shape, already covered above
        other_shapes[str(sid)] = text

    # Common voice correction: "change the title of slide 3 from X to Y".
    # Handle it deterministically so ASR/correction commands don't depend on
    # the editor LLM preserving the exact replacement text.
    title_to = re.search(r'\btitle\b.*?\bfrom\b.+?\bto\b\s+(.+)$', instruction, re.IGNORECASE)
    title_direct = re.search(r'\b(?:change|update|set|rename)\b.*?\btitle\b.*?\bto\b\s+(.+)$', instruction, re.IGNORECASE)
    replacement_title = (title_to or title_direct)
    if replacement_title:
        new_title = replacement_title.group(1).strip().strip('".,!?')
        if new_title:
            try:
                from backend.api.ppt import _patch_slide
                # other_edits intentionally omitted (defaults to None): a
                # title-only instruction must never touch bullets, notes, or
                # any other shape on the slide.
                def _writer(slide, _t=new_title, _b=old_bullets, _n=old_notes):
                    _patch_slide(slide, idx, _t, _b, _n, None)
                await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
                from backend.core.session_state import get_state
                state = get_state(session_id)
                state.last_ppt_action = {
                    "tool": "ppt_edit_slide",
                    "slide_number": idx + 1,
                    "instruction": instruction,
                    "changes": [f"title from '{old_title}' to '{new_title}'"],
                }
                return {"spoken_reply": f"I changed the title on slide {idx + 1} to {new_title}."}
            except Exception as e:
                logger.error(f"ppt_edit_slide title shortcut error: {e}")
                return {"spoken_reply": "I had trouble editing the slide right now."}

    # LLM call
    def _call() -> str:
        import ollama
        from backend.core.config import settings
        import json
        other_shapes_desc = (
            "\n".join(f'- shape_id {sid}: "{text}"' for sid, text in other_shapes.items())
            if other_shapes else "(none)"
        )
        prompt = f"""
You are an expert presentation editor. The user wants to edit a slide.
Current Slide:
Title: {old_title}
Bullets: {json.dumps(old_bullets)}
Notes: {old_notes}
Other text on this slide (subtitle/caption/footer — NOT the title or bullets):
{other_shapes_desc}

Instruction: {instruction}

Rules:
- Only change what the instruction actually asks for. Fields you are not asked
  to change must be echoed back EXACTLY as given above — do not paraphrase,
  shorten, or drop unrelated text.
- If the instruction targets one of the "other text" shapes above (a subtitle,
  caption, footer, etc.), put its new text in "other_edits" keyed by its
  shape_id. Do not invent shape_ids that weren't listed. Leave "other_edits"
  as {{}} if nothing else needs to change.

Output ONLY valid JSON matching this schema:
{{
  "title": "New Title",
  "bullets": ["New Bullet 1", "New Bullet 2"],
  "notes": "New Speaker Notes",
  "other_edits": {{"<shape_id>": "new text for that shape"}}
}}
"""
        resp = ollama.chat(
            model=settings.OLLAMA_MODEL,
            messages=[{"role": "user", "content": prompt}],
            # think=False is required alongside format="json" — without it, a
            # reasoning-capable model (Qwen3, etc.) tries to emit chain-of-thought
            # before its JSON, which conflicts with grammar-constrained JSON
            # decoding and can come back as empty content (see _summarize_ollama
            # above, and _ollama_call_template_content in
            # services/ppt_template_builder.py, for the same fix).
            options={"num_predict": 400},
            format="json",
            think=False,
            stream=False,
        )
        if isinstance(resp, dict):
            return resp["message"]["content"].strip()
        return resp.message.content.strip()

    try:
        raw = await asyncio.to_thread(_call)
        import json
        if not raw:
            # Same empty-content failure mode this fix targets — retry once
            # before giving up, since Ollama occasionally still returns a
            # blank message.strip() on a cold model load.
            raw = await asyncio.to_thread(_call)
        if not raw:
            logger.error("ppt_edit_slide error: Ollama returned an empty response after retry")
            return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}
        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            logger.error(f"ppt_edit_slide error: invalid JSON from Ollama: {raw[:300]!r}")
            return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}
        new_title = data.get("title", old_title)
        new_bullets = data.get("bullets", old_bullets)
        new_notes = data.get("notes", old_notes)

        # Only accept other_edits for shape_ids we actually listed to the
        # model — guards against a hallucinated id landing on the wrong shape.
        raw_other_edits = data.get("other_edits") or {}
        other_edits = {
            sid: text for sid, text in raw_other_edits.items()
            if sid in other_shapes and isinstance(text, str) and text.strip()
        }

        changes = []
        if new_title != old_title:
            changes.append(f"title from '{old_title}' to '{new_title}'")
        if new_bullets != old_bullets:
            changes.append("bullet content")
        if new_notes != old_notes:
            changes.append("speaker notes")
        for sid, text in other_edits.items():
            if text != other_shapes.get(sid):
                changes.append("other slide text")
                break
        
        from backend.api.ppt import _patch_slide
        def _writer(slide, _t=new_title, _b=new_bullets, _n=new_notes, _oe=other_edits or None):
            _patch_slide(slide, idx, _t, _b, _n, _oe)
        await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
        from backend.core.session_state import get_state
        state = get_state(session_id)
        state.last_ppt_action = {
            "tool": "ppt_edit_slide",
            "slide_number": idx + 1,
            "instruction": instruction,
            "changes": changes,
        }
        return {"spoken_reply": f"I've updated slide {idx + 1}."}
    except Exception as e:
        logger.error(f"ppt_edit_slide error: {e}")
        return {"spoken_reply": "I had trouble editing the slide right now."}


async def _edit_kind_aware_slide(effective_sid: str, idx: int, kind: str, instruction: str, session_id: str, source: str | None = None) -> dict:
    """
    Voice-edit path for slides cloned from the GD template (team, table,
    comparison, ...) — the generic title/numbered-bullets model above
    doesn't fit their structure at all. Same "send current state as JSON,
    ask the model for edited JSON back" pattern as the generic path, just
    with each kind's own field schema instead of title/bullets/notes.

    `source` identifies exactly which candidate slide (of possibly several
    for this kind — see _KIND_SOURCES) this one was originally cloned from;
    different candidates don't share shape_ids, so it's needed to compute
    the right slot map, not just the kind name.
    """
    from backend.services.ppt_template_builder import extract_slide_data, populate_slide_data, _parse_source
    import json

    parsed_source = _parse_source(source)

    def _read_current() -> dict:
        from pptx import Presentation
        prs = Presentation(f"data/ppt/{effective_sid}.pptx")
        return extract_slide_data(prs.slides[idx], kind, parsed_source)

    try:
        current_data = await asyncio.to_thread(_read_current)
    except Exception as e:
        logger.error(f"kind-aware edit: failed to read current slide data: {e}")
        return {"spoken_reply": "I had trouble reading that slide right now."}

    def _call() -> str:
        import ollama
        from backend.core.config import settings
        prompt = f"""
You are an expert presentation editor. The user wants to edit a "{kind}" slide.

Current content (JSON): {json.dumps(current_data)}

Instruction: {instruction}

Rules:
- Return the SAME JSON shape as "Current content" above, with the requested
  change applied.
- Only change what the instruction actually asks for — every other field
  must be echoed back EXACTLY as given, do not paraphrase or drop anything.
- Output ONLY valid JSON — no markdown, no explanation, nothing else.
"""
        resp = ollama.chat(
            model=settings.OLLAMA_MODEL,
            messages=[{"role": "user", "content": prompt}],
            options={"num_predict": 500},
            format="json",
            think=False,  # see ppt_edit_slide above for why this must pair with format="json"
            stream=False,
        )
        if isinstance(resp, dict):
            return resp["message"]["content"].strip()
        return resp.message.content.strip()

    try:
        raw = await asyncio.to_thread(_call)
        if not raw:
            raw = await asyncio.to_thread(_call)  # cold-model-load retry, same as the generic path
        if not raw:
            logger.error("kind-aware edit error: Ollama returned an empty response after retry")
            return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}
        try:
            new_data = json.loads(raw)
        except json.JSONDecodeError:
            logger.error(f"kind-aware edit error: invalid JSON from Ollama: {raw[:300]!r}")
            return {"spoken_reply": "I had trouble understanding that edit — could you try rephrasing it?"}

        def _writer(slide):
            populate_slide_data(slide, kind, new_data, parsed_source)

        await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
        from backend.core.session_state import get_state
        state = get_state(session_id)
        state.last_ppt_action = {
            "tool": "ppt_edit_slide",
            "slide_number": idx + 1,
            "instruction": instruction,
            "changes": ["slide content"],
        }
        return {"spoken_reply": f"I've updated slide {idx + 1}."}
    except Exception as e:
        logger.error(f"kind-aware ppt_edit_slide error: {e}")
        return {"spoken_reply": "I had trouble editing the slide right now."}


def _slide_bullets_for_notes(slide: dict) -> list:
    import re
    bullets = []
    for sh in slide.get("shapes", []):
        for line in sh.get("text", "").split("\n"):
            m = re.match(r'^\d+\.\s{1,3}(.+)', line)
            if m:
                bullets.append(m.group(1).strip())
    return bullets


def _generate_notes_text_sync(title: str, bullets: list) -> str:
    import ollama, json
    from backend.core.config import settings
    prompt = f"""
Write 3-4 natural conversational sentences of speaker notes for this slide.
Title: {title}
Bullets: {json.dumps(bullets)}

Output the speaker notes plainly without any markdown, prefix, or JSON formatting.
"""
    resp = ollama.chat(
        model=settings.OLLAMA_MODEL,
        messages=[{"role": "user", "content": prompt}],
        options={"num_predict": 250},
        stream=False,
    )
    if isinstance(resp, dict):
        return resp["message"]["content"].strip()
    return resp.message.content.strip()


async def ppt_generate_notes(args: dict, session_id: str) -> dict:
    from backend.api.ppt import (
        _slide_store, _latest_upload_sid, _current_slide,
        apply_notes_batch_async,
    )
    from backend.queues.bus import bus

    effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
    slides = _slide_store.get(effective_sid, [])
    if not slides:
        return {"spoken_reply": "No presentation is loaded."}

    if args.get("all"):
        await bus.emit_event("ppt_command", {"action": "goto", "index": 0}, session_id)
        notes_by_index: dict[int, str] = {}
        failures = 0
        for i, slide in enumerate(slides):
            try:
                notes_by_index[i] = await asyncio.to_thread(
                    _generate_notes_text_sync, slide.get("title", ""), _slide_bullets_for_notes(slide),
                )
            except Exception as e:
                failures += 1
                logger.error(f"ppt_generate_notes (all) slide {i} error: {e}")
        if not notes_by_index:
            return {"spoken_reply": "I had trouble generating notes right now."}
        await apply_notes_batch_async(effective_sid, notes_by_index)
        await bus.emit_event("ppt_command", {"action": "refresh", "session_id": effective_sid}, session_id)
        from backend.core.session_state import get_state
        state = get_state(session_id)
        state.last_ppt_action = {
            "tool": "ppt_generate_notes",
            "slide_number": None,
            "instruction": "generate speaker notes for all slides",
            "changes": [f"speaker notes for {len(notes_by_index)} of {len(slides)} slides"],
        }
        if failures:
            return {"spoken_reply": f"I've generated speaker notes for {len(notes_by_index)} of {len(slides)} slides — {failures} failed."}
        return {"spoken_reply": f"I've generated speaker notes for all {len(slides)} slides."}

    slide_number = args.get("slide_number")
    idx = int(slide_number) if slide_number is not None else _current_slide.get(effective_sid, 0)
    if idx < 0 or idx >= len(slides):
        return {"spoken_reply": "I'm not sure which slide to write notes for."}
    _current_slide[effective_sid] = idx
    await bus.emit_event("ppt_command", {"action": "goto", "index": idx}, session_id)
    slide = slides[idx]
    old_title = slide.get("title", "")
    old_bullets = _slide_bullets_for_notes(slide)

    try:
        notes = await asyncio.to_thread(_generate_notes_text_sync, old_title, old_bullets)
        # Apply notes, keep title and bullets the same
        from backend.api.ppt import _patch_slide
        def _writer(slide, _t=old_title, _b=old_bullets, _n=notes):
            _patch_slide(slide, idx, _t, _b, _n, None)
        await _apply_edit_and_background_refresh(effective_sid, idx, session_id, _writer)
        from backend.core.session_state import get_state
        state = get_state(session_id)
        state.last_ppt_action = {
            "tool": "ppt_generate_notes",
            "slide_number": idx + 1,
            "instruction": "generate speaker notes",
            "changes": ["speaker notes"],
        }
        return {"spoken_reply": f"I've generated new speaker notes for slide {idx + 1}."}
    except Exception as e:
        logger.error(f"ppt_generate_notes error: {e}")
        return {"spoken_reply": "I had trouble generating notes right now."}


async def ppt_add_slide(args: dict, session_id: str) -> dict:
    """Insert a new slide into the loaded presentation, content generated
    from a spoken description ("add a slide about our Q4 roadmap [after
    slide 3]"). If the instruction has no real topic ("add a slide" with
    nothing else), asks what it should be about instead of sending that
    bare trigger phrase to content generation — which doesn't fail on a
    topic-less prompt, it just invents plausible-sounding filler that reads
    as if the system "hallucinated" or echoed something already in the
    deck. The clarifying answer is picked up on the next turn via
    state.pending_add_slide (see services/front_llm.py's classify())."""
    from backend.api.ppt import _slide_store, _latest_upload_sid
    from backend.services.ppt_template_builder import generate_single_slide_content
    from backend.core.session_state import get_state

    effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
    slides = _slide_store.get(effective_sid, [])
    if not slides:
        return {"spoken_reply": "No presentation is loaded yet. Please upload or create one first."}

    state = get_state(session_id)
    pending = state.pending_add_slide
    instruction = args.get("instruction", "").strip()

    insert_after, position_specified = _parse_insert_position(instruction)
    # Whether the user ever named a position — this turn, or in the pending
    # clarification round. If they did, we honour it exactly; if they never
    # did, PILOT picks the best spot itself (from the slide titles) so the
    # deck keeps a coherent flow, and tells the user where it landed.
    if pending is not None:
        if position_specified:
            pass  # user gave a position on the clarification turn — use it
        elif pending.get("position_specified"):
            insert_after = pending.get("insert_after")
            position_specified = True
    user_chose_position = position_specified

    if not _has_real_topic(instruction):
        if pending is not None:
            # Already asked once this round and still got nothing to go on
            # — don't loop forever asking the same question.
            state.pending_add_slide = None
            return {"spoken_reply": "I still didn't catch what the slide should be about, so I'll leave it for now — just ask again whenever you're ready."}
        # Remember whether a position was already stated so we don't re-ask or
        # override it after we get the topic on the next turn.
        state.pending_add_slide = {"insert_after": insert_after, "position_specified": position_specified}
        where = "" if not position_specified else (
            "at the very beginning" if insert_after == -1
            else "at the end" if insert_after is None
            else f"right after slide {insert_after + 1}"
        )
        return {"spoken_reply": "Sure — what should the new slide be about?" + (f" I'll put it {where}." if where else "")}

    state.pending_add_slide = None

    try:
        slide_data = await generate_single_slide_content(instruction)
    except Exception as e:
        logger.error(f"ppt_add_slide generation error: {e}")
        slide_data = None

    if not slide_data:
        # Generation failed/timed out — still produce a real GD-template slide,
        # not a blank one: "text" is a valid template kind whose source slides
        # carry the Grid Dynamics branding.
        logger.warning("ppt_add_slide: generation returned nothing — using GD 'text' fallback")
        slide_data = {"kind": "text", "title": instruction[:60], "paragraphs": [instruction]}

    kind = slide_data.pop("kind", None)

    # GUARANTEE the GD template is used: every added slide must be one of the
    # known template kinds (each maps to real GD template source slides in
    # _KIND_SOURCES). If the model returned an unknown/blank kind, coerce to
    # "text" rather than letting add_slide_to_deck raise and — worse — ever
    # emit an off-brand slide.
    from backend.services.ppt_template_builder import _KIND_SOURCES
    if kind not in _KIND_SOURCES:
        logger.warning(f"ppt_add_slide: invalid kind {kind!r} — coercing to GD 'text'")
        title = slide_data.get("title") or instruction[:60]
        slide_data = {"title": title, "paragraphs": slide_data.get("paragraphs") or [instruction]}
        kind = "text"
    new_title = slide_data.get("title") or instruction[:60]

    # ── Auto-position: the user didn't say where, so read the existing slide
    # titles and let PILOT choose the spot that best preserves the narrative
    # flow, then announce it. Falls back to appending at the end. ──
    placement_note = ""
    if not user_chose_position:
        from backend.services.ppt_template_builder import pick_insert_position
        titles = [(s.get("title") or "") for s in slides]
        topic = new_title if kind not in ("agenda",) else (instruction or new_title)
        picked = await pick_insert_position(titles, topic)
        if picked is not None:
            insert_after, reason = picked
            if insert_after == -1:
                placement_note = " I placed it at the start"
            elif insert_after is None:
                placement_note = " I placed it at the end"
            else:
                placement_note = f" I placed it after slide {insert_after + 1}"
            placement_note += f" — {reason}." if reason else "."
        else:
            insert_after = None  # append at end
            placement_note = " I added it at the end."

    try:
        ok = await _apply_add_and_background_refresh(effective_sid, kind, slide_data, session_id, insert_after)
        if not ok:
            return {"spoken_reply": "I couldn't add a slide right now."}
        state.last_ppt_action = {
            "tool": "ppt_add_slide",
            "slide_number": len(slides) + 1,
            "instruction": instruction or "add a new slide",
            "changes": [f"new {kind} slide"],
        }
        return {"spoken_reply": f"I've added a new slide — {new_title}.{placement_note}"}
    except Exception as e:
        logger.error(f"ppt_add_slide error: {e}")
        return {"spoken_reply": "I had trouble adding that slide."}


async def _apply_add_and_background_refresh(effective_sid: str, kind: str, data: dict, session_id: str, insert_after: int | None = None) -> bool:
    """Same fast-reply-then-background-refresh pattern as
    _apply_edit_and_background_refresh, for the add-slide path: writes the
    new slide immediately, then regenerates thumbnails (LibreOffice, whole
    deck) in the background and re-emits ppt_command so the viewer catches
    up a moment after the spoken confirmation instead of before it."""
    from backend.api.ppt import add_slide_fast_async, refresh_slide_thumbnails_async
    from backend.queues.bus import bus

    new_index = await add_slide_fast_async(effective_sid, kind, data, insert_after)
    if new_index is None:
        return False

    async def _background_refresh():
        await refresh_slide_thumbnails_async(effective_sid)
        await bus.emit_event("ppt_command", {"action": "goto", "index": new_index}, session_id)
        await bus.emit_event("ppt_command", {"action": "refresh", "session_id": effective_sid}, session_id)

    asyncio.create_task(_background_refresh())
    return True


def _parse_reorder(text: str, total: int) -> tuple[int, int] | None:
    """Parse 'move slide X after/before slide Y', 'move slide X to position Y',
    'move slide X to the front/end', ordinal forms. Returns (from_index,
    to_index) both 0-based, or None if it can't find a clear source+target."""
    t = text.lower()

    def _num(word_or_digit: str) -> int | None:
        if word_or_digit.isdigit():
            return int(word_or_digit)
        return _ORDINALS.get(word_or_digit)

    # Source slide: "move slide 2" / "move the second slide" / "move slides 1 and 2"
    src = None
    m = re.search(r'\bslide[s]?\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
    if m:
        src = int(m.group(1))
    else:
        m = re.search(r'\b(' + "|".join(_ORDINALS) + r')\s+slide\b', t)
        if m:
            src = _ORDINALS[m.group(1)]
    if src is None or not (1 <= src <= total):
        return None
    from_index = src - 1

    # Target: after/before slide Y, to position Y, to front/end.
    # _move_slide removes the source, THEN inserts. So compute the target
    # against the POST-removal list: the anchor slide Y sits at index (y-1),
    # shifted down by one if the source was before it.
    def _anchor_index(y: int) -> int:
        idx = y - 1
        return idx - 1 if idx > from_index else idx

    m = re.search(r'\bafter\s+slide\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
    if m:
        to = _anchor_index(int(m.group(1))) + 1  # land right AFTER the anchor
        return from_index, max(0, min(to, total - 1))
    m = re.search(r'\bbefore\s+slide\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
    if m:
        to = _anchor_index(int(m.group(1)))       # land AT the anchor's slot (pushes it down)
        return from_index, max(0, min(to, total - 1))
    m = re.search(r'\b(?:to|at|into?|position)\s+(?:the\s+)?(?:slide|position)?\s*(?:number|no\.?|#)?\s*(\d+)\b', t)
    if m:
        return from_index, max(0, min(int(m.group(1)) - 1, total - 1))
    if re.search(r'\b(?:to the (?:front|beginning|start)|as the first)\b', t):
        return from_index, 0
    if re.search(r'\b(?:to the end|as the last)\b', t):
        return from_index, total - 1
    return None


async def ppt_reorder_slide(args: dict, session_id: str) -> dict:
    """Reorder a slide within the loaded deck ('move slide 1 after slide 3',
    'move slide 2 to position 4'). This is the real reorder capability — before
    it existed, such requests fell through to a generic answer that invented
    manual PowerPoint drag-and-drop steps."""
    from backend.api.ppt import _slide_store, _latest_upload_sid, reorder_slide_fast_async, refresh_slide_thumbnails_async
    from backend.queues.bus import bus

    effective_sid = session_id if session_id in _slide_store else _latest_upload_sid
    slides = _slide_store.get(effective_sid, [])
    total = len(slides)
    if total < 2:
        return {"spoken_reply": "There aren't enough slides to reorder yet."}

    instruction = args.get("instruction", "")
    parsed = _parse_reorder(instruction, total)
    if parsed is None:
        return {"spoken_reply": "Tell me which slide to move and where — for example, “move slide 1 after slide 3.”"}
    from_index, to_index = parsed
    if from_index == to_index:
        return {"spoken_reply": f"Slide {from_index + 1} is already in that position."}

    try:
        final_index = await reorder_slide_fast_async(effective_sid, from_index, to_index)
        if final_index is None:
            return {"spoken_reply": "No presentation is loaded to reorder."}

        async def _bg():
            await refresh_slide_thumbnails_async(effective_sid)
            await bus.emit_event("ppt_command", {"action": "goto", "index": final_index}, session_id)
            await bus.emit_event("ppt_command", {"action": "refresh", "session_id": effective_sid}, session_id)
        asyncio.create_task(_bg())

        # NOTE: upstream is missing this import at the call site below — every
        # other tool in this file imports get_state locally before using it;
        # without it this raises NameError the first time a reorder succeeds.
        from backend.core.session_state import get_state
        get_state(session_id).last_ppt_action = {
            "tool": "ppt_reorder_slide", "slide_number": final_index + 1,
            "instruction": instruction, "changes": [f"moved slide {from_index + 1} → position {final_index + 1}"],
        }
        return {"spoken_reply": f"Moved slide {from_index + 1} to position {final_index + 1}."}
    except ValueError as e:
        return {"spoken_reply": str(e)}
    except Exception as e:
        logger.error(f"ppt_reorder_slide error: {e}")
        return {"spoken_reply": "I had trouble reordering that slide."}
